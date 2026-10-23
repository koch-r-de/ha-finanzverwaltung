"""
Bildwerkzeug der Hilfe (seit v2.1).

Erzeugt die Bildschirmfotos für app/hilfe/bilder automatisch: legt in einem
temporären Ordner eine Demo-Datenbank mit erfundenen Beispieldaten an, startet
das Programm mit fester Uhrzeit (13.10.2026, 10:00 Uhr deutscher Zeit) und
fotografiert je Unterkapitel höchstens einen Ausschnitt per Playwright.
Der Dateiname ist der Anker des Unterkapitels, z. B. konten.anlegen.png.

Aufruf aus dem Wurzelordner des Repositorys:  python werkzeuge/hilfe_bilder.py
Voraussetzungen: die Pakete aus finanzverwaltung/requirements-docker.txt sowie playwright
(mit installiertem Chromium). Das Werkzeug gehört nicht ins Add-on-Image.
"""
import io
import os
import socket
import sys
import tempfile
import threading
import time
from datetime import datetime
from pathlib import Path

# Seit v2.5 liegt das Add-on im Unterordner finanzverwaltung/ des Repositorys.
WURZEL = Path(__file__).resolve().parent.parent / "finanzverwaltung"
BILDER = WURZEL / "app" / "hilfe" / "bilder"
FESTE_ZEIT = "2026-10-13T10:00:00"

DATEN = tempfile.mkdtemp(prefix="hilfe_bilder_")
os.environ["FINANZAPP_DATENVERZEICHNIS"] = DATEN
os.environ["FINANZAPP_FESTE_ZEIT"] = FESTE_ZEIT
os.environ.setdefault("FINANZAPP_VERSION", "Demo")
sys.path.insert(0, str(WURZEL))

from fastapi.testclient import TestClient  # noqa: E402
from app.main import app  # noqa: E402


def ok(r):
    assert r.status_code < 300, (r.status_code, r.text)
    return r.json()


def demo_daten_anlegen():
    """Erfundene Beispieldaten, keine echten Personen oder Konten."""
    with TestClient(app) as c:
        anna = ok(c.post("/api/personen", json={"name": "Anna", "farbe": "#4a7fd6"}))["id"]
        ben = ok(c.post("/api/personen", json={"name": "Ben", "farbe": "#c46210"}))["id"]
        konto = lambda **k: ok(c.post("/api/konten", json=k))["id"]
        giro = konto(name="Girokonto", typ="Girokonto", eroeffnungsdatum="2025-01-01", person_id=anna)
        giro_b = konto(name="Girokonto Ben", typ="Girokonto", eroeffnungsdatum="2025-01-01", person_id=ben)
        tg = konto(name="Tagesgeld", typ="Sparkonto", eroeffnungsdatum="2025-01-01")
        konto(name="Haushaltskasse", typ="Bargeld", eroeffnungsdatum="2025-01-01")
        kk = konto(name="Kreditkarte", typ="Kreditkarte", eroeffnungsdatum="2025-01-01", in_gesamtsaldo=False)
        kat = {k["name"]: k["id"] for k in ok(c.get("/api/kategorien"))}
        for k, d, s in [(giro, "2025-12-31", 3150), (giro_b, "2025-12-31", 1870), (tg, "2025-12-31", 8400)]:
            ok(c.post("/api/saldo-korrekturen", json={"konto_id": k, "datum": d, "saldo": s}))
        wk = lambda **v: ok(c.post("/api/wiederkehrend", json=v))
        wk(konto_id=giro, betrag=2650, kategorie_id=kat["Gehalt & Lohn"], beschreibung="Gehalt Anna", intervall="monatlich",
           tag_des_monats=28, startdatum="2026-01-01", person_id=anna)
        wk(konto_id=giro_b, betrag=2380, kategorie_id=kat["Gehalt & Lohn"], beschreibung="Gehalt Ben", intervall="monatlich",
           tag_des_monats=28, startdatum="2026-01-01", person_id=ben)
        wk(konto_id=giro, betrag=1120, kategorie_id=kat["Wohnen"], beschreibung="Miete", intervall="monatlich",
           tag_des_monats=1, startdatum="2026-01-01", inflationsflag=True)
        wk(konto_id=giro, betrag=74, kategorie_id=kat["Strom"], beschreibung="Abschlag Strom", intervall="monatlich",
           tag_des_monats=15, startdatum="2026-01-01")
        wk(konto_id=giro_b, betrag=39.99, kategorie_id=kat["Internet & Telefon"], beschreibung="Internet",
           intervall="monatlich", tag_des_monats=3, startdatum="2026-01-01", person_id=ben)
        wk(konto_id=giro, betrag=312, kategorie_id=kat["KFZ-Unterhalt"], beschreibung="Kfz-Versicherung",
           intervall="jaehrlich", tag_des_monats=1, monat=1, startdatum="2026-01-01")
        sz = ok(c.post("/api/budgets", json={"name": "Sommerurlaub", "typ": "sparziel", "konto_id": tg, "betrag": 150,
                                             "zielbetrag": 2400, "zeitraum": "monatlich", "startdatum": "2026-01-01"}))["id"]
        wk(konto_id=giro, zielkonto_id=tg, betrag=150, beschreibung="Sparrate Urlaub", intervall="monatlich",
           tag_des_monats=2, startdatum="2026-01-01", sparziel_id=sz)
        einkaeufe = [("2026-09-04", 63.18, "Lebensmittel/Verpflegung", "Wochenmarkt"), ("2026-09-11", 94.52, "Lebensmittel/Verpflegung", "Supermarkt"),
                     ("2026-09-19", 58.0, "Kraftstoff", "Tankstelle"), ("2026-09-24", 41.7, "Drogerie & Haushalt", "Drogerie"),
                     ("2026-10-02", 88.35, "Lebensmittel/Verpflegung", "Supermarkt"), ("2026-10-06", 61.4, "Kraftstoff", "Tankstelle"),
                     ("2026-10-09", 27.9, "Lebensmittel/Verpflegung", "Bäcker und Markt"), ("2026-10-12", 112.6, "Lebensmittel/Verpflegung", "Großeinkauf")]
        for d, b, k, t in einkaeufe:
            ok(c.post("/api/buchungen", json={"konto_id": giro, "datum": d, "betrag": b, "kategorie_id": kat[k], "typ": "Ist",
                                              "beschreibung": t}))
        ok(c.post("/api/buchungen", json={"konto_id": kk, "datum": "2026-10-07", "betrag": 189, "typ": "Ist",
                                          "kategorie_id": kat["Urlaub & Reisen"], "beschreibung": "Ferienwohnung Anzahlung"}))
        ok(c.post("/api/buchungen", json={"konto_id": giro, "datum": "2026-11-20", "betrag": 420, "typ": "Plan",
                                          "kategorie_id": kat["KFZ-Unterhalt"], "beschreibung": "Inspektion"}))
        ok(c.post("/api/budgets", json={"name": "Lebensmittel", "typ": "ausgabe", "kategorie_id": kat["Lebensmittel/Verpflegung"],
                                        "betrag": 450, "zeitraum": "monatlich", "startdatum": "2026-01-01", "warnungen": [{"schwelle": 80, "typ": "Warnung"}]}))
        ok(c.post("/api/budgets", json={"name": "Mobilität", "typ": "ausgabe", "kategorie_id": kat["Mobilität"], "betrag": 250,
                                        "zeitraum": "monatlich", "startdatum": "2026-01-01", "uebertrag": True,
                                        "uebertrag_modus": "kumuliert", "uebertrag_deckel": 1500}))
        ok(c.post("/api/budgets", json={"name": "Freizeit", "typ": "ausgabe", "kategorie_id": kat["Freizeit & Unterhaltung"],
                                        "betrag": 200, "zeitraum": "monatlich", "startdatum": "2026-01-01"}))
        dl = ok(c.post("/api/darlehen", json={"bezeichnung": "Wohnungskredit", "darlehenssumme": 180000, "zinssatz": 3.2,
                                              "zinsbeginn": "2025-04-01", "tilgungsbeginn": "2025-04-01", "laufzeit_jahre": 25,
                                              "ende_zinsbindung": "2035-03-31", "anschlusszinssatz": 4.0,
                                              "raten_modus": "feste_rate", "feste_rate": 870, "zahlungskonto_id": giro}))["id"]
        ok(c.post(f"/api/darlehen/{dl}/auszahlungen", json={"datum": "2025-04-01", "betrag": 180000}))
        ok(c.post(f"/api/darlehen/{dl}/sondertilgung", json={"konto_id": tg, "datum": "2026-03-16", "betrag": 5000}))
        ok(c.post("/api/altersvorsorge", json={"bezeichnung": "Private Rente", "anbieter": "Beispiel-Versicherung",
                                               "zahlungskonto_id": giro, "person_id": anna, "stand_datum": "2026-02-01",
                                               "hat_einzahlung": True, "beitrag_kategorie_id": kat["Altersvorsorge"],
                                               "beitrag_betrag": 100, "beitrag_intervall": "monatlich", "beitrag_tag_des_monats": 1,
                                               "beitrag_startdatum": "2026-01-01", "rentenbeginn": "2052-01-01",
                                               "auszahlungsart": "rente", "auszahlung_kategorie_id": kat["Rente"],
                                               "rente_garantiert_betrag": 280, "rente_voraussichtlich_betrag": 360}))
        ok(c.post("/api/altersvorsorge", json={"bezeichnung": "Betriebsrente", "zahlungskonto_id": giro_b, "person_id": ben,
                                               "stand_datum": "2026-05-01", "hat_einzahlung": False, "rentenbeginn": "2050-07-01",
                                               "auszahlungsart": "rente", "auszahlung_kategorie_id": kat["Rente"],
                                               "rente_garantiert_betrag": 190, "rente_voraussichtlich_betrag": 240}))
        ok(c.post("/api/regeln", json={"muster": "Tankstelle", "kategorie_id": kat["Kraftstoff"], "prioritaet": 100}))


def freier_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def server_starten(port: int):
    import uvicorn
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    for _ in range(100):
        if server.started:
            return server
        time.sleep(0.1)
    raise RuntimeError("Server startet nicht")


def speichern(png: bytes, anker: str):
    """PNG mit reduzierter Farbpalette speichern (kleinere Dateien)."""
    from PIL import Image
    bild = Image.open(io.BytesIO(png)).convert("RGB")
    bild = bild.quantize(colors=256, method=Image.Quantize.MEDIANCUT)
    bild.save(BILDER / f"{anker}.png", optimize=True)


def fotos(port: int):
    from playwright.sync_api import sync_playwright
    basis = f"http://127.0.0.1:{port}/"
    BILDER.mkdir(parents=True, exist_ok=True)
    zeit = datetime.fromisoformat(FESTE_ZEIT)
    with sync_playwright() as p:
        # Deutsche Datumsanzeige in Datumsfeldern braucht die Sprache auch in der Umgebung
        br = p.chromium.launch(args=["--lang=de-DE"], env={**os.environ, "LANG": "de_DE.UTF-8", "LANGUAGE": "de"})
        ctx = br.new_context(viewport={"width": 1280, "height": 800}, timezone_id="Europe/Berlin", locale="de-DE",
                             accept_downloads=True)
        page = ctx.new_page()
        page.clock.install(time=zeit)
        page.goto(basis)
        page.wait_for_timeout(1500)
        page.evaluate("Chart.defaults.animation = false")

        def gehe(ziel):
            page.click(f'.nav-eintrag[data-ziel="{ziel}"]')
            page.wait_for_timeout(1200)

        def block(anker, locator, max_hoehe=760):
            loc = page.locator(locator).first
            loc.scroll_into_view_if_needed()
            page.wait_for_timeout(300)
            box = loc.bounding_box()
            clip = {"x": box["x"] - 6, "y": box["y"] - 6, "width": box["width"] + 12,
                    "height": min(box["height"] + 12, max_hoehe)}
            speichern(page.screenshot(clip=clip), anker)

        def ansicht(anker):
            page.evaluate("document.querySelector('.hauptbereich').scrollTop = 0")
            page.wait_for_timeout(200)
            speichern(page.screenshot(), anker)

        gehe("start"); ansicht("start.uebersicht")
        gehe("konten"); block("konten.anlegen", "#form-konto >> xpath=..")
        gehe("buchungen"); block("buchungen.erfassen", "#form-buchung >> xpath=..")
        page.select_option("#buchungen-filter-herkunft", "anstehend"); page.wait_for_timeout(900)
        block("buchungen.liste", "#buchungen-tabelle >> xpath=../..", 620)
        gehe("wiederkehrend"); ansicht("wiederkehrend.anlegen")
        gehe("plan"); ansicht("plan.protokoll")
        gehe("budgets"); block("budgets.anlegen", "#form-budget >> xpath=../..")
        page.click('[data-budget-tab="uebersicht"]'); page.wait_for_timeout(900); ansicht("budgets.uebersicht")
        page.click('[data-budget-tab="details"]'); page.wait_for_timeout(1200)
        block("budgets.details", '[data-budget-tab-inhalt="details"] .block', 640)
        gehe("darlehen")
        page.locator("#darlehen-liste .karte").first.click(); page.wait_for_timeout(1200)
        block("darlehen.tilgungsplan", "#darlehen-detail-block", 560)
        gehe("altersvorsorge"); block("altersvorsorge.anlegen", "#form-altersvorsorge >> xpath=../..")
        gehe("auswertungen"); page.wait_for_timeout(800)
        block("auswertungen.kennzahlen", "#auswertung-kacheln")
        block("auswertungen.saldo", "#saldo-chart >> xpath=../..")
        block("auswertungen.cashflow", "#cashflow-chart >> xpath=../..")
        block("auswertungen.kategorien", "#kategorien-chart >> xpath=../..")
        # Mehrfachauswahl (seit v2.3): Kategorienbaum geöffnet, Wohnen ganz, Mobilität teilweise
        page.evaluate("document.querySelector('.hauptbereich').scrollTop = 0")
        page.click("#auswertung-kategorien .mehrfach-knopf"); page.wait_for_timeout(300)
        page.click('#auswertung-kategorien [data-aktion="keine"]')
        for name in ("Wohnen", "Kraftstoff"):
            page.locator("#auswertung-kategorien .mehrfach-zeile", has_text=name).first.locator("input").check()
        page.evaluate("""() => { const l = document.querySelector('#auswertung-kategorien .mehrfach-liste');
            const z = [...l.querySelectorAll('.mehrfach-zeile')].find((e) => e.textContent.trim() === 'Lebenshaltung');
            l.scrollTop = z.offsetTop - 8; }""")
        page.wait_for_timeout(300)
        filter_box = page.locator(".filter-block").bounding_box()
        liste_box = page.locator("#auswertung-kategorien .mehrfach-liste").bounding_box()
        unten = max(filter_box["y"] + filter_box["height"], liste_box["y"] + liste_box["height"])
        speichern(page.screenshot(clip={"x": filter_box["x"] - 6, "y": filter_box["y"] - 6,
                                        "width": filter_box["width"] + 12, "height": unten - filter_box["y"] + 12}),
                  "auswertungen.auswahl")
        page.click('#auswertung-kategorien [data-aktion="alle"]')
        page.keyboard.press("Escape"); page.wait_for_timeout(800)
        page.click('[data-auswertung-tab="darlehen"]'); page.wait_for_timeout(1500); ansicht("auswertungen.darlehen")
        page.click('[data-auswertung-tab="altersvorsorge"]'); page.wait_for_timeout(1500); ansicht("auswertungen.altersvorsorge")
        page.click('[data-auswertung-tab="allgemein"]')
        gehe("jahr"); ansicht("jahr.lesen")

        # Export ändern und als Vorschau hochladen
        gehe("admin")
        page.evaluate("document.getElementById('export-herunterladen').closest('details').open = true")
        with page.expect_download() as dl:
            page.click("#export-herunterladen")
        pfad = Path(DATEN) / "export.xlsx"
        dl.value.save_as(pfad)
        from openpyxl import load_workbook
        wb = load_workbook(pfad)
        ws = wb["Buchungen"]
        kopf = [c.value for c in ws[1]]
        sp = kopf.index("beschreibung") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, sp).value == "Großeinkauf":
                ws.cell(r, sp, "Großeinkauf Wochenende")
        ws.append([None, None] + [None] * (len(kopf) - 2))
        n = ws.max_row
        for name, wert in {"konto": "Haushaltskasse", "datum": "2026-10-12", "betrag": 12.5,
                           "kategorie": "Lebensmittel/Verpflegung", "beschreibung": "Wochenmarkt"}.items():
            ws.cell(n, kopf.index(name) + 1, wert)
        wb.save(pfad)
        page.set_input_files("#import-datei-feld", str(pfad))
        page.click("#import-pruefen"); page.wait_for_timeout(1500)
        block("admin.export-import", "#form-import >> xpath=../..", 700)

        # Wiederherstellen (seit v2.4): Backup anlegen, Fenster aus der Backup-Liste öffnen
        page.click('[data-admin-tab="system"]'); page.wait_for_timeout(600)
        page.evaluate("document.getElementById('backup-jetzt').closest('details').open = true")
        page.click("#backup-jetzt"); page.wait_for_timeout(800)
        page.locator("#backups-tabelle [data-wh-backup]").first.click(); page.wait_for_timeout(800)
        page.fill("#wh-phrase-eingabe", "")
        box = page.locator("#wh-dialog").bounding_box()
        speichern(page.screenshot(clip={"x": box["x"] - 6, "y": box["y"] - 6, "width": box["width"] + 12,
                                        "height": box["height"] + 12}), "admin.wiederherstellen")
        page.click("#wh-abbrechen"); page.wait_for_timeout(300)

        page.click('[data-admin-tab="zuruecksetzen"]'); page.wait_for_timeout(900)
        page.locator('.reset-art[value="plan_buchungen"]').check()
        page.locator('.reset-karte[data-art="plan_buchungen"] .hilfe-knopf').click()
        block("zuruecksetzen.ablauf", '[data-admin-tab-inhalt="zuruecksetzen"] .block', 700)

        gehe("hilfe"); page.wait_for_timeout(600)
        page.fill("#hilfe-suche", "Sparziel"); page.wait_for_timeout(600)
        ansicht("erste-schritte.bedienung")

        mobil = br.new_context(viewport={"width": 390, "height": 760}, timezone_id="Europe/Berlin", locale="de-DE",
                               device_scale_factor=1)
        m = mobil.new_page()
        m.clock.install(time=zeit)
        m.goto(basis + "schnellbuchung"); m.wait_for_timeout(1200)
        m.fill("#sb-betrag", "23,40"); m.fill("#sb-form [name=beschreibung]", "Wochenmarkt")
        speichern(m.screenshot(), "schnellbuchung.nutzen")
        br.close()


if __name__ == "__main__":
    demo_daten_anlegen()
    port = freier_port()
    server_starten(port)
    fotos(port)
    print("Bilder gespeichert in", BILDER)
