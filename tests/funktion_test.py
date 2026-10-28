"""Fachliche Tests v1.10 auf Basis der hochgezogenen 1.9.1-Datenbank."""
import os, sys, shutil, io, csv
from datetime import date, timedelta
from dateutil.relativedelta import relativedelta
quelle, ziel = sys.argv[2], sys.argv[3]
shutil.rmtree(ziel, ignore_errors=True); shutil.copytree(quelle, ziel)
os.environ["FINANZAPP_DATENVERZEICHNIS"] = ziel
sys.path.insert(0, sys.argv[1])
from fastapi.testclient import TestClient
from app.main import app

HEUTE = date.today()
ERG = []
def pruef(name, bed, info=""):
    ERG.append((name, bool(bed)))
    print(("OK   " if bed else "FEHL ") + name + (f"  [{info}]" if info and not bed else ""))

def ok(r):
    assert r.status_code < 300, (r.status_code, r.text)
    return r.json()

with TestClient(app) as c:
    K = {k["name"]: k["id"] for k in ok(c.get("/api/konten"))}
    KAT = {k["name"]: k["id"] for k in ok(c.get("/api/kategorien"))}
    UMB = next(k["id"] for k in ok(c.get("/api/kategorien")) if k["system_key"] == "umbuchung")
    giro, tg, giro2, kk = K["Girokonto"], K["Tagesgeld"], K["Girokonto Partnerin"], K["Kreditkarte"]
    von, bis = HEUTE.isoformat(), (HEUTE + relativedelta(years=1)).isoformat()

    # Referenzwerte vor der Umbuchungs-Vorlage
    cf0 = ok(c.get(f"/api/auswertungen/cashflow?von={von}&bis={bis}"))
    kat0 = ok(c.get(f"/api/auswertungen/kategorien?von={von}&bis={bis}"))
    jahr0 = ok(c.get("/api/auswertungen/jahresuebersicht"))
    budgets0 = {b["name"]: b for b in ok(c.get("/api/budgets"))}
    sv0 = ok(c.get(f"/api/auswertungen/saldo-verlauf?von={von}&bis={bis}"))

    # ── Punkt 3: Dynamik/Inflation ──
    r = c.post("/api/wiederkehrend", json={"konto_id": giro, "betrag": 30, "kategorie_id": KAT["Strom"], "intervall": "monatlich",
        "tag_des_monats": 20, "startdatum": "2027-01-01", "inflationsflag": True, "dynamik_aktiv": False})
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == r.json()["id"])
    pruef("P3 Inflation bleibt gesetzt, keine Scheindynamik", v["inflationsflag"] == 1 and v["dynamik_aktiv"] == 0 and v["dynamik_anzeige"] == "–")
    r = c.post("/api/wiederkehrend", json={"konto_id": giro, "betrag": 30, "kategorie_id": KAT["Strom"], "intervall": "monatlich",
        "tag_des_monats": 20, "startdatum": "2027-01-01", "dynamik_aktiv": True})
    pruef("P3 Dynamik ohne Prozent -> 422", r.status_code == 422, r.text)
    alt = next(x for x in ok(c.get("/api/wiederkehrend")) if x["beschreibung"] == "KFZ-Versicherung")
    r = c.put(f"/api/wiederkehrend/{alt['id']}", json={"konto_id": giro, "betrag": 480, "kategorie_id": alt["kategorie_id"],
        "intervall": "jaehrlich", "tag_des_monats": 1, "monat": 1, "startdatum": "2026-01-01", "beschreibung": "KFZ-Versicherung",
        "dynamik_aktiv": True, "dynamik_rhythmus_jahre": 1, "dynamik_prozent": 3, "dynamik_startjahr": 2027})
    pruef("P3 Alt-Dynamik mit Prozent bleibt speicherbar", r.status_code == 200, r.text)

    # ── Validierung Umbuchung ──
    basis = {"konto_id": giro, "betrag": 200, "intervall": "monatlich", "tag_des_monats": 5, "startdatum": "2026-08-01"}
    r = c.post("/api/wiederkehrend", json={**basis, "zielkonto_id": giro})
    pruef("Ziel = Quelle -> 422", r.status_code == 422, r.text)
    r = c.post("/api/wiederkehrend", json={**basis, "zielkonto_id": 9999})
    pruef("Ziel unbekannt -> 422", r.status_code == 422, r.text)
    r = c.post("/api/wiederkehrend", json={**basis, "kategorie_id": UMB})
    pruef("Buchung mit Kategorie Umbuchung -> 422", r.status_code == 422, r.text)
    r = c.post("/api/wiederkehrend", json={**basis})
    pruef("Buchung ohne Kategorie -> 422", r.status_code == 422, r.text)

    # ── Umbuchungs-Vorlage anlegen (mit Inflation/Dynamik/negativem Betrag -> erzwungen neutral) ──
    sz = ok(c.get("/api/sparziele"))[0]["id"]
    # Befund V (3.0): Salden vor der Vorlage merken. Andere Vorlagen (Gehalt am
    # 28.) buchen je nach Testtag schon beim Start, absolute Salden wären an den
    # Testtag gekoppelt.
    salden_vorher = {k["name"]: k["saldo"] for k in ok(c.get("/api/konten"))}
    sz_vorher = next(b for b in ok(c.get("/api/budgets")) if b["typ"] == "sparziel")["verbrauch"]
    vid = ok(c.post("/api/wiederkehrend", json={**basis, "betrag": -200, "zielkonto_id": tg, "kategorie_id": KAT["Wohnen"],
        "inflationsflag": True, "dynamik_aktiv": True, "dynamik_prozent": 5, "dynamik_rhythmus_jahre": 1,
        "beschreibung": "Sparrate", "sparziel_id": sz}))["id"]
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == vid)
    pruef("Vorlage normalisiert (Kategorie, Betrag, Infl., Dynamik)",
          v["kategorie_id"] == UMB and v["betrag"] == 200 and v["inflationsflag"] == 0 and v["dynamik_aktiv"] == 0
          and v["dynamik_prozent"] is None, v)
    pruef("Tabelle: Umbuchung → Zielkonto", v["kategorie_anzeige"] == "Umbuchung → Tagesgeld" and v["sparziel_name"], v["kategorie_anzeige"])
    erzeugt = [b for b in ok(c.get("/api/buchungen")) if b["wiederkehrend_id"] == vid]
    pruef("Generator: echte Übertragungen Aug/Sep/Okt", [b["datum"] for b in erzeugt] == ["2026-08-05", "2026-09-04", "2026-10-05"]
          and all(b["typ"] == "Uebertragung" and b["betrag"] == 200 and b["zielkonto_id"] == tg and b["kategorie_id"] == UMB
                  and b["sparziel_id"] == sz for b in erzeugt), [(b["datum"], b["typ"]) for b in erzeugt])
    prot = ok(c.get("/api/plan-buchungen/protokoll?limit=500"))
    pruef("Protokolleinträge vorhanden", all(any(p["buchung_id"] == b["id"] for p in prot) for b in erzeugt))
    salden = {k["name"]: k["saldo"] for k in ok(c.get("/api/konten"))}
    summe = 200 * len(erzeugt)
    pruef("Salden: Giro -200 und Tagesgeld +200 je erzeugter Übertragung",
          round(salden["Girokonto"] - salden_vorher["Girokonto"], 2) == -summe
          and round(salden["Tagesgeld"] - salden_vorher["Tagesgeld"], 2) == summe, (salden_vorher, salden))
    sz_nachher = next(b for b in ok(c.get("/api/budgets")) if b["typ"] == "sparziel")["verbrauch"]
    pruef("Sparziel-Ist +600", round(sz_nachher - sz_vorher, 2) == 600, (sz_vorher, sz_nachher))
    auto = ok(c.get("/api/buchungen?herkunft=wiederkehrend"))
    manu = ok(c.get("/api/buchungen?herkunft=manuell"))
    pruef("Herkunft: automatisch, nicht manuell", all(any(a["id"] == b["id"] for a in auto) for b in erzeugt)
          and not any(m["wiederkehrend_id"] == vid for m in manu))
    st = ok(c.get("/api/vorzeichen/konflikte"))
    pruef("Keine Vorzeichen-Konflikte", not st.get("vorlagen") and not st.get("buchungen"), st)

    # Neustart: keine Duplikate
    with TestClient(app) as c2:
        n = len([b for b in ok(c2.get("/api/buchungen")) if b["wiederkehrend_id"] == vid])
    pruef("Neustart: keine Duplikate", n == 3, n)

    # ── Vorschau Buchungsliste ──
    vors = [b for b in ok(c.get(f"/api/buchungen?inkl_anstehende=true&konto_id={tg}")) if b.get("virtuell") and b.get("wiederkehrend_id") == vid]
    pruef("Vorschau beim Zielkonto sichtbar (12 Termine)", len(vors) == 12 and all(b["typ"] == "Uebertragung" and b["zielkonto_id"] == tg
          and b["betrag"] == 200 for b in vors), len(vors))
    vors_p = [b for b in ok(c.get(f"/api/buchungen?inkl_anstehende=true&person_id={next(p['id'] for p in ok(c.get('/api/personen')) if p['name'] == 'Kim')}")) if b.get("virtuell") and b.get("wiederkehrend_id") == vid]
    pruef("Vorschau Personen-Filter über Quellkonto", len(vors_p) == 12, len(vors_p))

    # ── Auswertungen ──
    cf1 = ok(c.get(f"/api/auswertungen/cashflow?von={von}&bis={bis}"))
    diff_cf = [(m, round(b - a, 2), round(d - e, 2)) for m, a, b, e, d in
               zip(cf0["monate"], cf0["ausgaben"], cf1["ausgaben"], cf0["einnahmen"], cf1["einnahmen"])]
    pruef("Cashflow: Umbuchung ohne Effekt (nur Strom 30 € ab 2027)",
          all((dx == (30.0 if m >= "2027-01" else 0.0) or (m == bis[:7] and dx == 0.0)) and de == 0 for m, dx, de in diff_cf), diff_cf)
    kat1 = ok(c.get(f"/api/auswertungen/kategorien?von={von}&bis={bis}"))
    pruef("Kategorien-Diagramm ohne Umbuchung", not any(k["kategorie"] == "Umbuchung" for k in kat1))
    jahr1 = ok(c.get("/api/auswertungen/jahresuebersicht"))
    pruef("Jahresübersicht ohne Umbuchung", not any(k["kategorie"] == "Umbuchung" for k in jahr1.get("kategorien", [])))
    sv1 = ok(c.get(f"/api/auswertungen/saldo-verlauf?von={von}&bis={bis}"))
    s0 = {k["name"]: k["werte"] for k in sv0["konten"]}
    s1 = {k["name"]: k["werte"] for k in sv1["konten"]}
    # Differenz durch Umbuchung + Strom-Vorlage (ab 2027 nur Giro)
    d_tg = [round(b - a, 2) for a, b in zip(s0["Tagesgeld"], s1["Tagesgeld"])]
    d_gi = [round(b - a, 2) for a, b in zip(s0["Girokonto"], s1["Girokonto"])]
    print("   Diff Tagesgeld je Monat:", d_tg[:6])
    print("   Diff Giro je Monat:     ", d_gi[:6])
    erwartet_tg = [400 + 200 * i for i in range(len(d_tg))]  # Okt: Aug+Sep+Okt gebucht vorher 0? siehe unten
    pruef("Saldo-Verlauf Zielkonto steigt je Monat um 200", all(round(d_tg[i + 1] - d_tg[i], 2) == 200 for i in range(1, len(d_tg) - 1)), d_tg)
    pruef("Saldo-Verlauf Summe Giro+Tagesgeld (ohne Strom) neutral",
          all(round(d_tg[i] + d_gi[i], 2) in (0.0,) or HEUTE.year < 2027 for i in range(0, 3)), list(zip(d_tg, d_gi))[:3])

    # Budget auf Kategorie Umbuchung: keine Prognose aus Umbuchungen
    bu = ok(c.post("/api/budgets", json={"name": "UmbTest", "typ": "ausgabe", "kategorie_id": UMB, "betrag": 100,
        "zeitraum": "monatlich", "startdatum": "2026-01-01"}))["id"]
    b = ok(c.get(f"/api/budgets/{bu}"))
    v = ok(c.get(f"/api/budgets/details/verlauf?budget_id={bu}"))
    pruef("Budget auf Umbuchung: Verbrauch/Prognose 0", b["verbrauch"] == 0 and b["prognose"]["fix_offen"] == 0
          and not any(x for x in v["prognose"] if x), (b["verbrauch"], b["prognose"]["fix_offen"]))
    c.delete(f"/api/budgets/{bu}")

    # ── Befund B: Plan-Buchung 2027-01-20 (600 KFZ) in Saldo-Verlauf/Cashflow/Kategorien/Budget ──
    monate = sv1["monate"]
    i_jan = monate.index("2027-01-01"); i_dez = monate.index("2026-12-01")
    # Giro-Delta Dez->Jan enthält -600 Plan; Vergleich gegen Projektion ohne Plan: über cashflow
    cfj = dict(zip(cf1["monate"], cf1["ausgaben"]))
    pruef("B: Cashflow Jan 2027 enthält 600 € Plan-Buchung", cfj["2027-01"] >= 600, cfj["2027-01"])
    katj = ok(c.get("/api/auswertungen/kategorien?von=2027-01-01&bis=2027-01-31"))
    kfz = next((k for k in katj if k["kategorie"] == "KFZ-Unterhalt"), None)
    pruef("B: Kategorien-Diagramm enthält Plan-Buchung", kfz and kfz["summe"] >= 600, katj)
    mob = next(x for x in ok(c.get("/api/budgets")) if x["name"] == "Mobilität")
    vm = ok(c.get(f"/api/budgets/details/verlauf?budget_id={mob['id']}&von=2026-10&bis=2027-02"))
    pruef("B: Budget-Details Jan 2027 enthält Plan-Buchung", dict(zip(vm["monate"], vm["prognose"]))["2027-01"] >= 600, vm)

    # ── Befund D: Doppelzählung ──
    neu_k = ok(c.post("/api/konten", json={"name": "Testkonto D", "typ": "Girokonto", "eroeffnungsdatum": "2026-01-01"}))["id"]
    ok(c.post("/api/wiederkehrend", json={"konto_id": neu_k, "betrag": 100, "kategorie_id": KAT["Lebensmittel/Verpflegung"], "intervall": "monatlich",
        "tag_des_monats": 5, "startdatum": "2026-09-01"}))
    svd = ok(c.get(f"/api/auswertungen/saldo-verlauf?von={von}&bis={(HEUTE + relativedelta(months=3)).isoformat()}&konto_id={neu_k}"))
    pruef("D: kein Doppelzählen (Okt -200, Nov -300, Dez -400)", svd["konten"][0]["werte"][:3] == [-200.0, -300.0, -400.0], svd["konten"][0]["werte"])

    # ── Befund F: Termin morgen in fix_offen ──
    morgen = HEUTE + timedelta(days=1)
    while morgen.weekday() >= 5:  # Wochenende würde auf heute zurückverschoben
        morgen += timedelta(days=1)
    ok(c.post("/api/wiederkehrend", json={"konto_id": neu_k, "betrag": 30, "kategorie_id": KAT["Drogerie & Haushalt"], "intervall": "monatlich",
        "tag_des_monats": morgen.day, "startdatum": morgen.isoformat(), "enddatum": morgen.isoformat()}))
    bf = ok(c.post("/api/budgets", json={"name": "Drogerie", "typ": "ausgabe", "kategorie_id": KAT["Drogerie & Haushalt"], "betrag": 100,
        "zeitraum": "monatlich", "startdatum": "2026-01-01"}))["id"]
    pruef("F: Termin morgen in Fix-Prognose", ok(c.get(f"/api/budgets/{bf}"))["prognose"]["fix_offen"] == 30,
          ok(c.get(f"/api/budgets/{bf}"))["prognose"])

    # ── Plan-Protokoll: Anpassen einer erzeugten Übertragung ──
    letzte = erzeugt[-1]
    neu_datum = (HEUTE + timedelta(days=20)).isoformat()
    ok(c.post(f"/api/plan-buchungen/{letzte['id']}/ablehnen-anpassen?neues_datum={neu_datum}"))
    b = ok(c.get(f"/api/buchungen/{letzte['id']}"))
    pruef("Anpassen: Typ bleibt Übertragung", b["typ"] == "Uebertragung" and b["datum"] == neu_datum and b["zielkonto_id"] == tg, b)
    salden2 = {k["name"]: k["saldo"] for k in ok(c.get("/api/konten"))}
    pruef("Anpassen: künftige Übertragung wirkt heute noch nicht", salden2["Tagesgeld"] == salden["Tagesgeld"] - 200, salden2)
    with TestClient(app) as c2:
        n = len([x for x in ok(c2.get("/api/buchungen")) if x["wiederkehrend_id"] == vid])
    pruef("Anpassen + Neustart: kein Duplikat", n == 3, n)
    ue = ok(c.get("/api/admin/zuruecksetzen/uebersicht"))
    vs = ok(c.post("/api/admin/zuruecksetzen/vorschau", json={"arten": ["plan_buchungen"]}))
    pruef("Reset-Vorschau Plan-Buchungen zählt angepasste Übertragung", vs["zeilen"][0]["anzahl"] == 2
          and vs["zeilen"][0]["zusatz"][0]["anzahl"] == 1, vs["zeilen"])

    # ── Typwechsel Umbuchung -> Buchung ──
    v2 = ok(c.post("/api/wiederkehrend", json={"konto_id": giro2, "zielkonto_id": giro, "betrag": 50, "intervall": "monatlich",
        "tag_des_monats": 1, "startdatum": "2027-01-01", "beschreibung": "Haushaltsgeld"}))["id"]
    r = c.put(f"/api/wiederkehrend/{v2}", json={"konto_id": giro2, "betrag": 50, "kategorie_id": KAT["Lebensmittel/Verpflegung"], "intervall": "monatlich",
        "tag_des_monats": 1, "startdatum": "2027-01-01", "beschreibung": "Haushaltsgeld", "zielkonto_id": None, "sparziel_id": sz})
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == v2)
    pruef("Typwechsel -> Buchung: Ziel/Sparziel weg, Vorzeichen negativ", r.status_code == 200 and v["zielkonto_id"] is None
          and v["sparziel_id"] is None and v["betrag"] == -50, v)
    r = c.put(f"/api/wiederkehrend/{v2}", json={"konto_id": giro2, "betrag": 50, "intervall": "monatlich", "zielkonto_id": giro,
        "tag_des_monats": 1, "startdatum": "2027-01-01", "beschreibung": "Haushaltsgeld"})
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == v2)
    pruef("Typwechsel -> Umbuchung", r.status_code == 200 and v["zielkonto_id"] == giro and v["betrag"] == 50, v)

    # ── Konten leeren Hinweis zählt Zielkonto-Vorlagen ──
    vs = ok(c.post("/api/admin/zuruecksetzen/vorschau", json={"arten": ["konten_leeren"], "konten_leeren_ids": [tg]}))
    pruef("Konten leeren: Hinweis nennt Umbuchungs-Vorlage", any("aktive wiederkehrende Vorlage" in w for w in vs["warnungen"]), vs["warnungen"])

    # ── Konten löschen Teilauswahl: Tagesgeld -> Giro ──
    vs = ok(c.post("/api/admin/zuruecksetzen/vorschau", json={"arten": ["konten_loeschen"], "konten_loeschen_ids": [tg], "ziel_konto_id": giro}))
    z = {x["label"]: x["anzahl"] for x in vs["zeilen"][0]["zusatz"]}
    print("   Vorschau-Zusatz:", z)
    pruef("Konten löschen: Vorschau nennt deaktivierte Umbuchungen", any("deaktiviert" in k and n == 1 for k, n in z.items()), z)
    tok = ok(c.post("/api/admin/zuruecksetzen/vorschau", json={"arten": ["konten_loeschen"], "konten_loeschen_ids": [tg], "ziel_konto_id": giro}))
    # Phrase holen und ausführen
    import app.core.zuruecksetzen as zr
    norm_req = {"arten": ["konten_loeschen"], "konten_loeschen_ids": [tg], "ziel_konto_id": giro}
    phrase_info = tok.get("token"), tok.get("phrase")
    r = c.post("/api/admin/zuruecksetzen/ausfuehren", json={**norm_req, "token": phrase_info[0], "phrase": phrase_info[1]})
    pruef("Konten löschen ausgeführt", r.status_code == 200, r.text[:300])
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == vid)
    pruef("Umbuchung Giro->Tagesgeld jetzt Giro->Giro und deaktiviert", v["zielkonto_id"] == giro and v["aktiv"] == 0, v)
    v = next(x for x in ok(c.get("/api/wiederkehrend")) if x["id"] == v2)
    pruef("Umbuchung Partnerin->Giro bleibt aktiv", v["aktiv"] == 1 and v["zielkonto_id"] == giro, v)
    r = c.put(f"/api/wiederkehrend/{vid}/aktiv?aktiv=true")
    pruef("Reaktivieren bei Ziel=Quelle -> 422", r.status_code == 422, r.text)

    # ── Import ──
    xl = c.get("/api/admin/import-vorlage")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xl.content))
    kopf = [x.value for x in wb["Wiederkehrend"][1]]
    pruef("xlsx-Vorlage hat id/aktion/zielkonto/sparziel, keine Daten", {"id", "aktion", "zielkonto", "sparziel"} <= set(kopf) and wb["Wiederkehrend"].max_row == 1, kopf)
    ok(c.post("/api/konten", json={"name": "Tagesgeld neu", "typ": "Sparkonto", "eroeffnungsdatum": "2026-01-01"}))
    ok(c.post("/api/konten", json={"name": "Girokonto Zweit", "typ": "Girokonto", "eroeffnungsdatum": "2026-01-01"}))
    zeilen = [
        ["konto", "betrag", "kategorie", "zielkonto", "sparziel", "beschreibung", "intervall", "tag_des_monats", "startdatum", "inflationsflag", "dynamik_aktiv"],
        ["Girokonto", "100", "", "Tagesgeld neu", "", "OK Umbuchung", "monatlich", "10", "2027-01-01", "0", "0"],
        ["Girokonto", "100", "Umbuchung", "Tagesgeld neu", "Urlaub", "OK Umbuchung mit Kat+Sparziel", "monatlich", "10", "2027-01-01", "", ""],
        ["Giro", "100", "", "Tagesgeld neu", "", "Konto mehrdeutig", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "", "neu", "", "Ziel per Teiltreffer", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "Wohnen", "Tagesgeld neu", "", "Kategorie-Widerspruch", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "", "Tagesgeld neu", "", "Inflation-Widerspruch", "monatlich", "10", "2027-01-01", "1", ""],
        ["Girokonto", "100", "Wohnen", "", "Urlaub", "Sparziel ohne Ziel", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "Umbuchung", "", "", "Kat Umbuchung ohne Ziel", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "", "Girokonto", "", "Ziel gleich Quelle", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "", "Gibtsnicht", "", "Ziel unbekannt", "monatlich", "10", "2027-01-01", "", ""],
        ["Girokonto", "100", "Wohnen", "", "", "OK normale Buchung", "monatlich", "10", "2027-01-01", "", ""],
    ]
    def csv_vorschau(zz):
        buf = io.StringIO(); csv.writer(buf).writerows(zz)
        return ok(c.post("/api/admin/import/vorschau", data={"art": "wiederkehrend"},
                         files={"datei": ("t.csv", buf.getvalue().encode("utf-8"), "text/csv")}))
    r = csv_vorschau(zeilen)
    fehler_zeilen = {z["zeile"]: z["hinweis"] for z in r["bereiche"][0]["fehler"]}
    erg = {zz[5]: (i + 2 not in fehler_zeilen, fehler_zeilen.get(i + 2, "ok")) for i, zz in enumerate(zeilen[1:])}
    for k, v in erg.items(): print("   ", k, v)
    erwartet_ok = {"OK Umbuchung", "OK Umbuchung mit Kat+Sparziel", "Ziel per Teiltreffer", "OK normale Buchung"}
    pruef("Import: genau die erwarteten Zeilen ok", {k for k, v in erg.items() if v[0]} == erwartet_ok, erg)
    pruef("Import: bei Fehlern kein Token (alles oder nichts)", r["token"] is None and r["fehler"] == 7, r["fehler"])
    pruef("Import: Mehrdeutigkeit nennt Kandidaten", "mehrdeutig" in erg["Konto mehrdeutig"][1])
    r = csv_vorschau([zeilen[0]] + [zz for zz in zeilen[1:] if zz[5] in erwartet_ok])
    pruef("Import: fehlerfreie Datei -> Token, 4 neu", r["token"] and r["bereiche"][0]["neu"] == 4, r)
    ok(c.post("/api/admin/import/ausfuehren", json={"token": r["token"]}))
    imp = [x for x in ok(c.get("/api/wiederkehrend")) if x["beschreibung"] in ("OK Umbuchung", "OK Umbuchung mit Kat+Sparziel")]
    pruef("Import: Umbuchungs-Vorlagen korrekt", all(x["zielkonto_id"] == K_neu for x in imp) if (K_neu := next(k["id"] for k in ok(c.get("/api/konten")) if k["name"] == "Tagesgeld neu")) else False, imp)

    # Budget löschen mit Sparziel-Referenz in Vorlage
    sz_imp = next(x for x in imp if x["sparziel_id"])
    r = c.delete(f"/api/budgets/{sz_imp['sparziel_id']}")
    pruef("Sparziel löschen trotz Vorlagen-Referenz", r.status_code == 200, r.text)

fehl = [n for n, o in ERG if not o]
print(f"\n{len(ERG) - len(fehl)}/{len(ERG)} bestanden", fehl)
