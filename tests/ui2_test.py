"""Oberflächentest der 2.0-Funktionen (Playwright). Server muss auf 8935 laufen."""
import io, re, sys
from datetime import date
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
BASIS = "http://localhost:8935/"
import os
SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ausgabe", "bilder", "ui2") + "/"
import os; os.makedirs(SHOTS, exist_ok=True)
ERG = []
def pruef(n, b, info=""):
    ERG.append((n, bool(b))); print(("OK   " if b else "FEHL ") + n + (f"  [{info}]" if info and not b else ""))

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(timezone_id="Europe/Berlin", viewport={"width": 1400, "height": 1000}, accept_downloads=True)
    page = ctx.new_page()
    fehler, anfragen, dialoge = [], [], []
    page.on("console", lambda m: fehler.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: fehler.append(str(e)))
    page.on("request", lambda r: anfragen.append(r.url))
    def dialog(d):
        dialoge.append(d.message); d.accept()
    page.on("dialog", dialog)
    page.goto(BASIS); page.wait_for_timeout(1500)
    heute = page.evaluate("heuteISO()")

    def gehe(ziel):
        page.click(f'.nav-eintrag[data-ziel="{ziel}"]'); page.wait_for_timeout(900)
    def schalter(id_):
        anfragen.clear(); page.check(f"#{id_}"); page.wait_for_timeout(1200)
        mit = [u for u in anfragen if "inflation=true" in u]
        anfragen.clear(); page.uncheck(f"#{id_}"); page.wait_for_timeout(1000)
        ohne = [u for u in anfragen if "/api/" in u and "inflation=true" not in u]
        return mit, ohne

    # Start: Sparziel-Banner
    banner = page.locator("#sparziel-banner").inner_text()
    pruef("Start: Banner meldet unstimmige Sparziel-Zuordnung", "Sparziel" in banner, banner)
    page.screenshot(path=SHOTS + "start_banner.png")

    # Auswertungen: Cashflow mit anstehendem Anteil, Schalter
    gehe("auswertungen"); page.wait_for_timeout(1200)
    # seit 2.3: Ausgaben nach unten in einem gemeinsamen Stapel, dazu die Linie „Überschuss“
    ds = page.evaluate("CHART_CASHFLOW.data.datasets.filter(d => d.type === 'bar').map(d => [d.label, d.stack, d.data.reduce((a,b)=>a+Math.abs(b||0),0)])")
    pruef("Cashflow: 4 Säulen-Datensätze, gestapelt (gebucht + anstehend)", len(ds) == 4 and len({d[1] for d in ds}) == 1, ds)
    pruef("Cashflow: anstehender Anteil vorhanden", any("anstehend" in (d[0] or "").lower() and d[2] > 0 for d in ds), ds)
    pruef("Inflation-Schalter Auswertungen standardmäßig aus", not page.is_checked("#auswertung-inflation"))
    mit, ohne = schalter("auswertung-inflation")
    pruef("Auswertungen: Schalter lädt mit inflation=true", any("saldo-verlauf" in u for u in mit) and any("cashflow" in u for u in mit), mit)
    pruef("Auswertungen: zurückschalten lädt ohne Inflation", any("cashflow" in u for u in ohne), ohne)
    page.click('[data-auswertung-tab="altersvorsorge"]'); page.wait_for_timeout(900)
    anfragen.clear(); page.check("#auswertung-inflation"); page.wait_for_timeout(1200)
    pruef("Auswertungen/Altersvorsorge: Schalter wirkt", any("inflation=true" in u and "altersvorsorge" in u for u in anfragen), anfragen)
    page.uncheck("#auswertung-inflation"); page.wait_for_timeout(600)
    page.click('[data-auswertung-tab="allgemein"]'); page.wait_for_timeout(900)
    page.screenshot(path=SHOTS + "auswertungen.png", full_page=True)

    # Jahresübersicht, Budgets, Buchungsliste
    gehe("jahr")
    mit, _ = schalter("jahr-inflation")
    pruef("Jahresübersicht: Schalter wirkt", any("jahresuebersicht" in u for u in mit), mit)
    pruef("Jahresübersicht: Hinweis auf anstehende Termine", "anstehend" in page.locator("#abschnitt-jahr").inner_text())
    gehe("budgets")
    mit, _ = schalter("budgets-inflation-schalter")
    pruef("Budgets: Schalter wirkt", any("/budgets" in u for u in mit), mit)
    gehe("buchungen")
    page.select_option("#buchungen-filter-herkunft", "anstehend"); page.wait_for_timeout(700)
    mit, _ = schalter("buchungen-inflation")
    pruef("Buchungsliste: Schalter wirkt", any("/api/buchungen" in u for u in mit), mit)
    page.select_option("#buchungen-filter-herkunft", ""); page.wait_for_timeout(500)

    # Buchungsformular: Sparziel nur bei Übertragung mit Trägerkonto
    f = "#form-buchung"
    page.select_option("#buchung-typ", "Ist"); page.wait_for_timeout(200)
    pruef("Buchung Ist: kein Sparziel-Feld", not page.locator("#sparziel-feld").is_visible())
    page.select_option("#buchung-typ", "Uebertragung"); page.wait_for_timeout(200)
    page.select_option(f + ' [name="konto_id"]', label="Girokonto")
    page.select_option(f + ' [name="zielkonto_id"]', label="Girokonto Partnerin"); page.wait_for_timeout(200)
    opts_ohne = page.locator(f + ' [name="sparziel_id"] option').all_inner_texts()
    page.select_option(f + ' [name="zielkonto_id"]', label="Tagesgeld"); page.wait_for_timeout(200)
    opts_mit = page.locator(f + ' [name="sparziel_id"] option').all_inner_texts()
    pruef("Sparziel-Auswahl nur bei Trägerkonto als Quelle/Ziel",
          not any("Urlaub" in o for o in opts_ohne) and any("Urlaub" in o for o in opts_mit), (opts_ohne, opts_mit))
    page.select_option("#buchung-typ", "Ist")

    # Saldo-Korrektur: kein Datum in der Zukunft
    gehe("konten")
    mx = page.locator('#form-saldo-korrektur [name="datum"]').get_attribute("max")
    pruef("Saldo-Korrektur: max = heute", mx == heute, mx)

    # Altersvorsorge: Rentenbeginn verschieben -> Bestätigungsdialog
    konten = page.request.get(BASIS + "api/konten").json()
    giro = next(k["id"] for k in konten if k["name"] == "Girokonto")
    kats = page.request.get(BASIS + "api/kategorien").json()
    def kat(n): return next(k["id"] for k in kats if k["name"] == n)
    r = page.request.post(BASIS + "api/altersvorsorge", data={
        "bezeichnung": "UI-AV", "zahlungskonto_id": giro, "stand_datum": "2026-01-01", "hat_einzahlung": True,
        "beitrag_kategorie_id": kat("Altersvorsorge") if any(k["name"] == "Altersvorsorge" for k in kats) else kats[0]["id"],
        "beitrag_betrag": 50, "beitrag_intervall": "monatlich", "beitrag_tag_des_monats": 15, "beitrag_startdatum": "2026-01-01",
        "rentenbeginn": "2027-01-01", "auszahlungsart": "rente",
        "auszahlung_kategorie_id": kat("Rente") if any(k["name"] == "Rente" for k in kats) else kats[0]["id"],
        "rente_garantiert_betrag": 300})
    pruef("AV angelegt", r.ok, r.text())
    gehe("altersvorsorge")
    page.locator("#altersvorsorge-tabelle tr", has_text="UI-AV").locator("button", has_text="Ändern").click()
    page.wait_for_timeout(500)
    page.fill('#form-altersvorsorge [name="rentenbeginn"]', "2026-08-01")
    falsch = sum(1 for b in page.request.get(BASIS + "api/buchungen").json()
                 if b.get("beschreibung", "").startswith("UI-AV") and b["betrag"] < 0 and b["datum"] >= "2026-08-01")
    dialoge.clear()
    page.click("#av-submit"); page.wait_for_timeout(1500)
    pruef("AV: Bestätigungsdialog nennt Anzahl", dialoge and f"{falsch} automatisch erzeugte" in dialoge[0], (falsch, dialoge))
    bs = [b for b in page.request.get(BASIS + "api/buchungen").json() if b.get("beschreibung", "").startswith("UI-AV")]
    beitr = sorted(b["datum"] for b in bs if b["betrag"] < 0); renten = sorted(b["datum"] for b in bs if b["betrag"] > 0)
    pruef("AV: Beiträge ab Aug gelöscht, Renten Aug–Okt gebucht", len(beitr) == 7 and len(renten) == 3, (beitr, renten))

    # Administration: Sparziele prüfen und lösen
    gehe("admin")
    page.evaluate("document.getElementById('admin-sparziele-block').open = true")
    page.click("#admin-sparziele-pruefen"); page.wait_for_timeout(800)
    erg = page.locator("#admin-sparziele-ergebnis").inner_text()
    pruef("Sparziele prüfen: Konflikt gelistet", "UI-Konflikt" in erg, erg)
    page.locator("#admin-sparziele-ergebnis input[type=checkbox]").first.check()
    page.click("#admin-sparziele-loesen"); page.wait_for_timeout(1000)
    k = page.request.get(BASIS + "api/sparziele/konflikte").json()
    pruef("Sparziele: Zuordnung gelöst", not k["buchungen"] and not k["vorlagen"], k)

    # Export / Import
    page.evaluate("document.getElementById('export-herunterladen').closest('details').open = true")
    with page.expect_download() as dl:
        page.click("#export-herunterladen")
    pfad = SHOTS + dl.value.suggested_filename; dl.value.save_as(pfad)
    pruef("Export: Dateiname xlsx", dl.value.suggested_filename.endswith(".xlsx"), dl.value.suggested_filename)
    page.set_input_files("#import-datei-feld", pfad)
    page.click("#import-pruefen"); page.wait_for_timeout(1500)
    txt = page.locator("#import-ergebnis").inner_text()
    pruef("Import unverändert: keine Änderungen", "Keine Änderungen" in txt, txt[:300])
    wb = load_workbook(pfad); ws = wb["Buchungen"]
    kopf = [c.value for c in ws[1]]
    zeile = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, kopf.index("beschreibung") + 1).value == "UI-Konflikt")
    ws.cell(zeile, kopf.index("beschreibung") + 1, "UI-Konflikt geändert")
    neu = SHOTS + "export_geaendert.xlsx"; wb.save(neu)
    page.set_input_files("#import-datei-feld", neu)
    page.click("#import-pruefen"); page.wait_for_timeout(1500)
    txt = page.locator("#import-ergebnis").inner_text()
    pruef("Import Vorschau: 1 Änderung bereit", "1 Änderung" in txt, txt[:300])
    page.screenshot(path=SHOTS + "import_vorschau.png", full_page=True)
    page.click("#import-uebernehmen"); page.wait_for_timeout(2000)
    bs = page.request.get(BASIS + "api/buchungen").json()
    pruef("Import übernommen", any(b["beschreibung"] == "UI-Konflikt geändert" for b in bs))

    # System: Serverzeit
    page.click('[data-admin-tab="system"]'); page.wait_for_timeout(800)
    z = page.locator("#systeminfo-zeit").inner_text()
    pruef("Serverzeit mit Europe/Berlin angezeigt", "Europe/Berlin" in z, z)
    page.screenshot(path=SHOTS + "system.png")
    pruef("Keine Konsolenfehler", not fehler, fehler)
    ctx.close(); br.close()
fehl = [n for n, o in ERG if not o]
print(f"\n{len(ERG) - len(fehl)}/{len(ERG)} bestanden", fehl)
