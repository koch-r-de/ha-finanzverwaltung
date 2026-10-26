"""
XLSX-Roundtrip (seit v2.0): Export aller Bereiche in eine xlsx-Datei,
Bearbeiten in Excel/LibreOffice, Re-Upload mit Vorschau und Übernahme.
Ersetzt den bisherigen reinen Neuanlage-Import (import_export.py liefert
nur noch die Wert-Konvertierung); alte Vorlagen-Dateien ohne Spalte „id“
funktionieren weiter und legen ausschließlich neu an.

Festlegungen aus dem Sparring (v2.0):
  - Umfang: Personen, Konten, Kategorien, Regeln, Budget (inkl. Sparziele),
    Darlehen, Darlehen-Auszahlungen, Altersvorsorge, Wiederkehrend,
    Buchungen (Export-Zeitraum wählbar, Standard alle).
  - Spalte „id“: leer = neu, sonst Änderung des Eintrags. Spalte „aktion“ =
    „löschen“ löscht. Fehlende Zeilen bleiben unverändert. Konten lassen
    sich per Datei nicht löschen (dafür „Daten zurücksetzen“).
  - Kategorien stehen als vollständiger Pfad („Ausgaben > Wohnen > Strom“),
    beim Upload gilt der Pfad oder ein eindeutiger Name. Konten, Personen,
    Sparziele und Darlehen werden über den Namen eindeutig aufgelöst.
  - Verborgene Spalte „_pruefwert“ je Zeile: wurde ein Eintrag nach dem
    Export in der App geändert, meldet die Vorschau einen Konflikt.
  - Ablauf: Vorschau -> alles oder nichts. Gibt es einen Fehler, wird
    nichts übernommen. Sonst Bestätigung, automatisches Backup, Übernahme
    in EINER Transaktion.

Technik: Vorschau und Übernahme laufen durch denselben Code. Die Vorschau
führt alles in einer Transaktion aus (jede Zeile in einem Savepoint, damit
Fehler je Zeile gesammelt werden) und rollt am Ende zurück - so prüft sie
exakt das, was die Übernahme tun wird, inklusive aller Folgeprüfungen der
API (Sparziel-Regel, Umbuchungen, Phasen der Altersvorsorge ...). Die
Fachfunktionen committen intern; während des Roundtrips sind Commits über
eine Hülle (_OhneCommit) wirkungslos.
"""
import csv
import hashlib
import sqlite3
import io
import json
import secrets
import time
from datetime import date, datetime
from typing import Optional

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from . import main as api
from .core import backup, zeit
from .core import feiertage as feiertage_modul
from .core import buchungen as bm
from .core import budgets as budgets_modul
from .core import personen as personen_modul
from .core import wiederkehrend as wk_modul
from .core import altersvorsorge as av_modul
from .core.db import kategorie_id as system_kategorie_id
from .core.import_export import _bool01, _datum, _ganz, _text_dekodieren, _txt, _zahl

PFAD_TRENNER = " > "
LOESCHEN = {"löschen", "loeschen", "delete"}
_TOKEN_GUELTIG_SEKUNDEN = 600
_TOKENS: dict[str, dict] = {}


class ZeilenFehler(Exception):
    """Fachlicher Fehler einer einzelnen Zeile (Meldung für die Vorschau)."""


class _OhneCommit:
    """Hülle um die Verbindung: commit() wirkungslos, rollback() verboten -
    der Roundtrip steuert die Transaktion selbst (alles oder nichts)."""

    def __init__(self, conn):
        self._conn = conn

    def commit(self):
        pass

    def rollback(self):
        raise RuntimeError("Interner Fehler: Rollback innerhalb des Imports.")

    def __getattr__(self, name):
        return getattr(self._conn, name)


# ─── Feld-Arten ──────────────────────────────────────────────────────

class Feld:
    def __init__(self, name, art="text", spalte=None, pflicht=False, standard=None, nur_lesen=False,
                 hinweis=None):
        self.name = name
        self.art = art            # text, zahl, ganz, datum, bool, konto, kategorie, person, sparziel, darlehen
        self.spalte = spalte or name
        self.pflicht = pflicht
        self.standard = standard
        self.nur_lesen = nur_lesen
        self.hinweis = hinweis


def _norm_zahl(v):
    return None if v is None else round(float(v), 6)


class Kontext:
    """Namensauflösung gegen den LIVE-Stand der Datenbank (neu angelegte
    Einträge desselben Uploads sind sofort auffindbar) mit Rückgriff auf
    die Namen zum Zeitpunkt des Uploads - so bleiben exportierte Zeilen
    gültig, wenn derselbe Upload z. B. eine Kategorie umbenennt."""

    def __init__(self, conn):
        self.conn = conn
        self._cache = {}
        self.ursprung = {
            "konto": self._namen("SELECT id, name FROM konten"),
            "person": self._namen("SELECT id, name FROM personen"),
            "sparziel": self._namen("SELECT id, name FROM budgets WHERE typ = 'sparziel'"),
            "darlehen": self._namen("SELECT id, bezeichnung AS name FROM darlehen"),
            "kategorie": self._pfade(),
        }

    def neu_laden(self):
        self._cache = {}

    def _namen(self, sql):
        return [(r["id"], r["name"]) for r in self.conn.execute(sql).fetchall()]

    def _pfade(self):
        rows = {r["id"]: r for r in self.conn.execute("SELECT id, name, parent_id FROM kategorien").fetchall()}

        def pfad(kid, tiefe=0):
            r = rows.get(kid)
            if r is None or tiefe > 30:
                return "?"
            return r["name"] if r["parent_id"] is None else pfad(r["parent_id"], tiefe + 1) + PFAD_TRENNER + r["name"]
        return [(kid, pfad(kid)) for kid in rows]

    def liste(self, art):
        if art not in self._cache:
            self._cache[art] = {
                "konto": lambda: self._namen("SELECT id, name FROM konten"),
                "person": lambda: self._namen("SELECT id, name FROM personen"),
                "sparziel": lambda: self._namen("SELECT id, name FROM budgets WHERE typ = 'sparziel'"),
                "darlehen": lambda: self._namen("SELECT id, bezeichnung AS name FROM darlehen"),
                "kategorie": self._pfade,
            }[art]()
        return self._cache[art]

    def name(self, art, id_):
        if id_ is None:
            return None
        return next((n for i, n in self.liste(art) if i == id_), None)

    def aufloesen(self, art, text) -> Optional[int]:
        t = _txt(text)
        if not t:
            return None
        bezeichnung = {"konto": "Konto", "person": "Person", "sparziel": "Sparziel", "darlehen": "Darlehen",
                       "kategorie": "Kategorie"}[art]
        for quelle in (self.liste(art), self.ursprung[art]):
            treffer = self._suchen(art, quelle, t)
            if treffer is not None:
                return treffer
        raise ZeilenFehler(f"{bezeichnung} „{t}“ nicht gefunden")

    def _suchen(self, art, quelle, t):
        tl = t.lower()
        if art == "kategorie":
            exakt = [i for i, p in quelle if p.lower() == tl]
            if len(exakt) == 1:
                return exakt[0]
            # sonst eindeutiger Name (letztes Pfadglied), dann eindeutiger Teiltreffer
            namen = [(i, p.split(PFAD_TRENNER)[-1]) for i, p in quelle]
        else:
            namen = quelle
        exakt = [(i, n) for i, n in namen if n.lower() == tl]
        if len(exakt) == 1:
            return exakt[0][0]
        kandidaten = exakt or [(i, n) for i, n in namen if tl in n.lower()]
        if len(kandidaten) == 1:
            return kandidaten[0][0]
        if len(kandidaten) > 1:
            if art == "kategorie":
                kandidaten = [(i, p) for i, p in quelle if i in {k[0] for k in kandidaten}]
            raise ZeilenFehler(f"„{t}“ ist mehrdeutig: " + ", ".join(sorted(n for _, n in kandidaten)))
        return None


def _wert_aus_datei(ctx: Kontext, feld: Feld, roh):
    """Zellwert -> normalisierter Wert (Referenzen als ID)."""
    if feld.art in ("konto", "kategorie", "person", "sparziel", "darlehen"):
        return ctx.aufloesen(feld.art, roh)
    try:
        if feld.art == "text":
            return _txt(roh) or None
        if feld.art == "zahl":
            return _norm_zahl(_zahl(roh))
        if feld.art == "ganz":
            return _ganz(roh)
        if feld.art == "datum":
            d = _datum(roh)
            if d and not 1900 <= int(d[:4]) <= 9999:
                raise ValueError
            return d
        if feld.art == "bool":
            t = _txt(roh)
            return feld.standard if t == "" else int(_bool01(roh))
    except (ValueError, TypeError):
        raise ZeilenFehler(f"Spalte „{feld.name}“: ungültiger Wert „{_txt(roh)}“")
    raise ValueError(feld.art)


def _anzeige(ctx: Kontext, feld: Feld, wert):
    if wert is None:
        return None
    if feld.art in ("konto", "person", "sparziel", "darlehen", "kategorie"):
        return ctx.name(feld.art, wert)
    if feld.art == "datum":
        return date.fromisoformat(wert)
    return wert


# ─── Bereiche ────────────────────────────────────────────────────────

class Bereich:
    key = ""
    sheet = ""
    tabelle = ""
    felder: list[Feld] = []
    loeschen_erlaubt = True
    titel_feld = "name"

    def export_zeilen(self, conn, params):
        return conn.execute(f"SELECT * FROM {self.tabelle} ORDER BY id").fetchall()

    def aus_db(self, conn, row) -> dict:
        werte = {}
        for f in self.felder:
            if f.nur_lesen:
                continue
            v = row[f.spalte]
            if f.art == "zahl":
                v = _norm_zahl(v)
            elif f.art == "bool":
                v = int(bool(v)) if v is not None else f.standard
            werte[f.name] = v
        return werte

    def nur_lesen_werte(self, conn, ctx, row) -> dict:
        return {}

    def zeile_lesen(self, conn, row):
        return conn.execute(f"SELECT * FROM {self.tabelle} WHERE id = ?", (row,)).fetchone()

    def anlegen(self, conn, ctx, werte) -> int:
        raise NotImplementedError

    def aendern(self, conn, ctx, id_, werte, alt) -> None:
        raise NotImplementedError

    def loeschen(self, conn, ctx, id_, alt) -> None:
        raise NotImplementedError

    def hinweise_vor_aenderung(self, conn, ctx, id_, werte, alt) -> list[str]:
        return []


def _pflicht(werte, *namen):
    fehlend = [n for n in namen if werte.get(n) in (None, "")]
    if fehlend:
        raise ZeilenFehler("Pflichtfeld(er) fehlen/ungültig: " + ", ".join(fehlend))


def _modell(klasse, **kw):
    try:
        return klasse(**kw)
    except ValidationError as e:
        teile = []
        for err in e.errors():
            ort = ".".join(str(x) for x in err.get("loc", []))
            teile.append(f"{ort}: {err.get('msg')}")
        raise ZeilenFehler("; ".join(teile))


class Personen(Bereich):
    key, sheet, tabelle = "personen", "Personen", "personen"
    felder = [Feld("name", pflicht=True), Feld("farbe"), Feld("aktiv", "bool", standard=1)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM personen ORDER BY name").fetchall()

    def anlegen(self, conn, ctx, w):
        _pflicht(w, "name")
        pid = personen_modul.person_anlegen(conn, w["name"], w.get("farbe"))
        if w.get("aktiv") == 0:
            personen_modul.person_bearbeiten(conn, pid, aktiv=0)
        ctx.neu_laden()
        return pid

    def aendern(self, conn, ctx, id_, w, alt):
        _pflicht(w, "name")
        personen_modul.person_bearbeiten(conn, id_, name=w["name"], farbe=w.get("farbe"), aktiv=w.get("aktiv", 1))
        ctx.neu_laden()

    def loeschen(self, conn, ctx, id_, alt):
        # wie in der Oberfläche: deaktivieren statt löschen
        personen_modul.person_loeschen(conn, id_)


class Konten(Bereich):
    key, sheet, tabelle = "konten", "Konten", "konten"
    loeschen_erlaubt = False
    felder = [Feld("name", pflicht=True),
              Feld("typ", pflicht=True, hinweis="Girokonto, Sparkonto, Kreditkarte oder Bargeld"),
              Feld("eroeffnungsdatum", "datum", pflicht=True), Feld("person", "person", "person_id"),
              Feld("in_gesamtsaldo", "bool", standard=1), Feld("farbe"), Feld("aktiv", "bool", standard=1),
              Feld("saldo_heute", "zahl", nur_lesen=True)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM konten ORDER BY name").fetchall()

    def nur_lesen_werte(self, conn, ctx, row):
        return {"saldo_heute": bm.saldo(conn, row["id"])}

    def _pruefen(self, w):
        _pflicht(w, "name", "typ", "eroeffnungsdatum")
        if w["typ"] not in ("Girokonto", "Sparkonto", "Kreditkarte", "Bargeld"):
            raise ZeilenFehler("typ muss Girokonto, Sparkonto, Kreditkarte oder Bargeld sein")

    def anlegen(self, conn, ctx, w):
        self._pruefen(w)
        kid = bm.konto_anlegen(conn, w["name"], w["typ"], w["eroeffnungsdatum"], w.get("person"),
                               bool(w.get("in_gesamtsaldo", 1)), w.get("farbe"))
        if w.get("aktiv") == 0:
            bm.konto_bearbeiten(conn, kid, aktiv=0)
        ctx.neu_laden()
        return kid

    def aendern(self, conn, ctx, id_, w, alt):
        self._pruefen(w)
        bm.konto_bearbeiten(conn, id_, name=w["name"], typ=w["typ"], eroeffnungsdatum=w["eroeffnungsdatum"],
                            person_id=w.get("person"), in_gesamtsaldo=w.get("in_gesamtsaldo", 1),
                            farbe=w.get("farbe"), aktiv=w.get("aktiv", 1))
        ctx.neu_laden()


class Kategorien(Bereich):
    key, sheet, tabelle = "kategorien", "Kategorien", "kategorien"
    felder = [Feld("pfad", nur_lesen=True), Feld("name", pflicht=True),
              Feld("oberkategorie", "kategorie", "parent_id", hinweis="Pfad der Oberkategorie, leer = oberste Ebene"),
              Feld("siri_favorit", "bool", standard=0),
              Feld("typ", nur_lesen=True), Feld("system", nur_lesen=True)]

    def export_zeilen(self, conn, params):
        ctx = Kontext(conn)
        pfade = dict(ctx.liste("kategorie"))
        rows = conn.execute("SELECT * FROM kategorien").fetchall()
        return sorted(rows, key=lambda r: pfade.get(r["id"], "").lower())

    def nur_lesen_werte(self, conn, ctx, row):
        return {"pfad": ctx.name("kategorie", row["id"]), "typ": row["typ"] or "neutral",
                "system": "ja" if row["system_key"] else ""}

    def anlegen(self, conn, ctx, w):
        _pflicht(w, "name")
        kid = bm.kategorie_anlegen(conn, w["name"], w.get("oberkategorie"))
        if w.get("siri_favorit"):
            bm.siri_favorit_setzen(conn, kid, True)
        ctx.neu_laden()
        return kid

    def aendern(self, conn, ctx, id_, w, alt):
        _pflicht(w, "name")
        if alt["system_key"] and (w["name"] != alt["name"] or w.get("oberkategorie") != alt["parent_id"]):
            raise ZeilenFehler("Diese Kategorie ist fest vorgegeben; Name und Oberkategorie sind nicht änderbar")
        eltern = w.get("oberkategorie")
        gesehen = set()
        while eltern is not None and eltern not in gesehen:
            if eltern == id_:
                raise ZeilenFehler("Die Oberkategorie liegt unterhalb dieser Kategorie (Zyklus)")
            gesehen.add(eltern)
            r = conn.execute("SELECT parent_id FROM kategorien WHERE id = ?", (eltern,)).fetchone()
            eltern = r["parent_id"] if r else None
        if alt["system_key"]:
            bm.siri_favorit_setzen(conn, id_, bool(w.get("siri_favorit")))
        else:
            try:
                bm.kategorie_bearbeiten(conn, id_, w["name"], w.get("oberkategorie"),
                                        siri_favorit=bool(w.get("siri_favorit")))
            except ValueError as e:
                raise ZeilenFehler(str(e))
        ctx.neu_laden()

    def loeschen(self, conn, ctx, id_, alt):
        if alt["system_key"]:
            raise ZeilenFehler("Diese Kategorie ist fest vorgegeben und kann nicht gelöscht werden")
        if conn.execute("SELECT 1 FROM kategorien WHERE parent_id = ? LIMIT 1", (id_,)).fetchone():
            raise ZeilenFehler("Kategorie hat Unterkategorien")
        for sql, was in (("SELECT COUNT(*) FROM buchungen WHERE kategorie_id = ?", "Buchungen"),
                         ("SELECT COUNT(*) FROM wiederkehrende_buchungen WHERE kategorie_id = ?", "Vorlagen"),
                         ("SELECT COUNT(*) FROM budgets WHERE kategorie_id = ?", "Budgets"),
                         ("SELECT COUNT(*) FROM altersvorsorge WHERE beitrag_kategorie_id = ? OR auszahlung_kategorie_id = ?",
                          "Altersvorsorge-Verträgen")):
            n = conn.execute(sql, (id_,) * sql.count("?")).fetchone()[0]
            if n:
                raise ZeilenFehler(f"Kategorie wird noch verwendet ({n} {was})")
        conn.execute("DELETE FROM kategorien WHERE id = ?", (id_,))
        ctx.neu_laden()


class Regeln(Bereich):
    key, sheet, tabelle = "regeln", "Regeln", "kategorie_regeln"
    titel_feld = "muster"
    felder = [Feld("muster", pflicht=True), Feld("kategorie", "kategorie", "kategorie_id", pflicht=True),
              Feld("prioritaet", "ganz", standard=100), Feld("aktiv", "bool", standard=1)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM kategorie_regeln ORDER BY prioritaet, id").fetchall()

    def anlegen(self, conn, ctx, w):
        _pflicht(w, "muster", "kategorie")
        cur = conn.execute("INSERT INTO kategorie_regeln (muster, kategorie_id, prioritaet, aktiv) VALUES (?, ?, ?, ?)",
                           (w["muster"], w["kategorie"], w.get("prioritaet") or 100, w.get("aktiv", 1)))
        return cur.lastrowid

    def aendern(self, conn, ctx, id_, w, alt):
        _pflicht(w, "muster", "kategorie")
        conn.execute("UPDATE kategorie_regeln SET muster = ?, kategorie_id = ?, prioritaet = ?, aktiv = ? WHERE id = ?",
                     (w["muster"], w["kategorie"], w.get("prioritaet") or 100, w.get("aktiv", 1), id_))

    def loeschen(self, conn, ctx, id_, alt):
        conn.execute("DELETE FROM kategorie_regeln WHERE id = ?", (id_,))


class Budgets(Bereich):
    key, sheet, tabelle = "budget", "Budget", "budgets"
    felder = [Feld("name", pflicht=True), Feld("typ", standard="ausgabe", hinweis="ausgabe, einnahme oder sparziel"),
              Feld("kategorie", "kategorie", "kategorie_id"), Feld("inkl_unterkategorien", "bool", standard=1),
              Feld("konto", "konto", "konto_id", hinweis="nur bei Sparziel: Trägerkonto"),
              Feld("zielbetrag", "zahl"), Feld("zieldatum", "datum"), Feld("betrag", "zahl", pflicht=True),
              Feld("zeitraum", standard="monatlich", hinweis="monatlich, quartalsweise, jaehrlich oder frei"),
              Feld("intervall_wert", "ganz"), Feld("intervall_einheit", hinweis="tage, wochen oder monate"),
              Feld("startdatum", "datum", pflicht=True), Feld("enddatum", "datum"), Feld("uebertrag", "bool", standard=0),
              Feld("uebertrag_modus", hinweis="letzte oder kumuliert"), Feld("uebertrag_deckel", "zahl"),
              Feld("inflationsflag", "bool", standard=0), Feld("person", "person", "person_id"),
              Feld("aktiv", "bool", standard=1)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM budgets ORDER BY typ, name").fetchall()

    def aus_db(self, conn, row):
        w = super().aus_db(conn, row)
        if w.get("zeitraum") == "benutzerdefiniert":
            w["zeitraum"] = "frei"
        return w

    def _payload(self, w, warnungen=None):
        _pflicht(w, "name", "betrag", "startdatum")
        typ = (w.get("typ") or "ausgabe").lower()
        if typ not in ("ausgabe", "einnahme", "sparziel"):
            raise ZeilenFehler("typ muss ausgabe, einnahme oder sparziel sein")
        zeitraum = (w.get("zeitraum") or "monatlich").lower()
        zeitraum = "benutzerdefiniert" if zeitraum in ("frei", "benutzerdefiniert") else zeitraum
        if zeitraum not in ("monatlich", "quartalsweise", "jaehrlich", "benutzerdefiniert"):
            raise ZeilenFehler("zeitraum muss monatlich, quartalsweise, jaehrlich oder frei sein")
        if typ == "sparziel" and not w.get("konto"):
            raise ZeilenFehler("Ein Sparziel braucht ein Trägerkonto (Spalte konto)")
        return _modell(api.BudgetIn, name=w["name"], typ=typ,
                       kategorie_id=None if typ == "sparziel" else w.get("kategorie"),
                       inkl_unterkategorien=bool(w.get("inkl_unterkategorien", 1)),
                       konto_id=w.get("konto") if typ == "sparziel" else None,
                       zielbetrag=w.get("zielbetrag"), zieldatum=w.get("zieldatum"), betrag=w["betrag"],
                       zeitraum=zeitraum, intervall_wert=w.get("intervall_wert"),
                       intervall_einheit=w.get("intervall_einheit"), startdatum=w["startdatum"],
                       enddatum=w.get("enddatum"), uebertrag=bool(w.get("uebertrag")),
                       uebertrag_modus=w.get("uebertrag_modus"), uebertrag_deckel=w.get("uebertrag_deckel"),
                       inflationsflag=bool(w.get("inflationsflag")), warnungen=warnungen,
                       person_id=w.get("person"))

    def anlegen(self, conn, ctx, w):
        bid = api._budget_anlegen(conn, self._payload(w))
        if w.get("aktiv") == 0:
            conn.execute("UPDATE budgets SET aktiv = 0 WHERE id = ?", (bid,))
        ctx.neu_laden()
        return bid

    def aendern(self, conn, ctx, id_, w, alt):
        # vorhandene, individuell gepflegte Warnschwellen erhalten
        warnungen = budgets_modul.warnungen_lesen(alt) or None
        api._budget_aendern(conn, id_, self._payload(w, warnungen))
        conn.execute("UPDATE budgets SET aktiv = ? WHERE id = ?", (w.get("aktiv", 1), id_))
        ctx.neu_laden()

    def loeschen(self, conn, ctx, id_, alt):
        api._budget_loeschen(conn, id_)
        ctx.neu_laden()


class Darlehen(Bereich):
    key, sheet, tabelle = "darlehen", "Darlehen", "darlehen"
    titel_feld = "bezeichnung"
    felder = [Feld("bezeichnung", pflicht=True), Feld("darlehenssumme", "zahl"), Feld("zinssatz", "zahl", pflicht=True),
              Feld("zinsbeginn", "datum", pflicht=True), Feld("tilgungsbeginn", "datum", pflicht=True),
              Feld("laufzeit_jahre", "ganz", pflicht=True), Feld("ende_zinsbindung", "datum", pflicht=True),
              Feld("anschlusszinssatz", "zahl", pflicht=True),
              Feld("raten_modus", pflicht=True, hinweis="tilgungssatz oder feste_rate"),
              Feld("tilgungssatz", "zahl"), Feld("feste_rate", "zahl"), Feld("anschluss_feste_rate", "zahl"),
              Feld("zahlungskonto", "konto", "zahlungskonto_id"), Feld("person", "person", "person_id")]

    def _payload(self, w):
        _pflicht(w, "bezeichnung", "zinssatz", "zinsbeginn", "tilgungsbeginn", "laufzeit_jahre",
                 "ende_zinsbindung", "anschlusszinssatz", "raten_modus")
        if w["raten_modus"] not in ("tilgungssatz", "feste_rate"):
            raise ZeilenFehler("raten_modus muss tilgungssatz oder feste_rate sein")
        return _modell(api.DarlehenIn, bezeichnung=w["bezeichnung"], zinssatz=w["zinssatz"],
                       zinsbeginn=w["zinsbeginn"], tilgungsbeginn=w["tilgungsbeginn"],
                       laufzeit_jahre=w["laufzeit_jahre"], ende_zinsbindung=w["ende_zinsbindung"],
                       anschlusszinssatz=w["anschlusszinssatz"], raten_modus=w["raten_modus"],
                       darlehenssumme=w.get("darlehenssumme"), tilgungssatz=w.get("tilgungssatz"),
                       feste_rate=w.get("feste_rate"), anschluss_feste_rate=w.get("anschluss_feste_rate"),
                       zahlungskonto_id=w.get("zahlungskonto"), person_id=w.get("person"))

    def anlegen(self, conn, ctx, w):
        did = api._darlehen_anlegen(conn, self._payload(w))
        ctx.neu_laden()
        return did

    def aendern(self, conn, ctx, id_, w, alt):
        api._darlehen_aendern(conn, id_, self._payload(w))
        ctx.neu_laden()

    def loeschen(self, conn, ctx, id_, alt):
        api._darlehen_loeschen(conn, id_)
        ctx.neu_laden()

    def hinweise_vor_aenderung(self, conn, ctx, id_, w, alt):
        if alt is not None and id_ is not None and w is None:
            n = conn.execute("SELECT COUNT(*) FROM buchungen WHERE ziel_darlehen_id = ?", (id_,)).fetchone()[0]
            return [f"Darlehen „{alt['bezeichnung']}“ gelöscht: {n} Buchung(en) bleiben bestehen und werden gelöst."]
        return []


class Auszahlungen(Bereich):
    key, sheet, tabelle = "auszahlungen", "Darlehen-Auszahlungen", "darlehen_auszahlungen"
    titel_feld = "darlehen"
    felder = [Feld("darlehen", "darlehen", "darlehen_id", pflicht=True, hinweis="Bezeichnung des Darlehens"),
              Feld("datum", "datum", pflicht=True), Feld("betrag", "zahl", pflicht=True)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM darlehen_auszahlungen ORDER BY darlehen_id, datum").fetchall()

    def anlegen(self, conn, ctx, w):
        _pflicht(w, "darlehen", "datum", "betrag")
        api._auszahlung_anlegen(conn, w["darlehen"], _modell(api.AuszahlungIn, datum=w["datum"], betrag=w["betrag"]))
        return conn.execute("SELECT MAX(id) FROM darlehen_auszahlungen").fetchone()[0]

    def aendern(self, conn, ctx, id_, w, alt):
        _pflicht(w, "darlehen", "datum", "betrag")
        if w["darlehen"] != alt["darlehen_id"]:
            raise ZeilenFehler("Das Darlehen einer Auszahlung lässt sich nicht ändern (löschen und neu anlegen)")
        api._auszahlung_aendern(conn, id_, w["datum"], w["betrag"])

    def loeschen(self, conn, ctx, id_, alt):
        api._auszahlung_loeschen(conn, id_)

    def hinweise_vor_aenderung(self, conn, ctx, id_, w, alt):
        if id_ is not None:
            return ["Auszahlung geändert/gelöscht: bereits gebuchte Zins- und Tilgungsraten bleiben unverändert, "
                    "nur der Tilgungsplan rechnet ab sofort mit dem neuen Stand."]
        return []


_AV_TEXT = ["bezeichnung", "anbieter", "beitrag_intervall", "auszahlungsart"]


class Altersvorsorge(Bereich):
    key, sheet, tabelle = "altersvorsorge", "Altersvorsorge", "altersvorsorge"
    titel_feld = "bezeichnung"
    felder = [Feld("bezeichnung", pflicht=True), Feld("anbieter"),
              Feld("zahlungskonto", "konto", "zahlungskonto_id", pflicht=True), Feld("person", "person", "person_id"),
              Feld("stand_datum", "datum", pflicht=True), Feld("hat_einzahlung", "bool", standard=1),
              Feld("beitrag_kategorie", "kategorie", "beitrag_kategorie_id"), Feld("beitrag_betrag", "zahl"),
              Feld("beitrag_intervall", standard="monatlich"), Feld("beitrag_tag_des_monats", "ganz", standard=1),
              Feld("beitrag_monat", "ganz"), Feld("beitrag_startdatum", "datum"),
              Feld("beitrag_dynamik_aktiv", "bool", standard=0), Feld("beitrag_dynamik_rhythmus_jahre", "ganz"),
              Feld("beitrag_dynamik_prozent", "zahl"), Feld("beitrag_dynamik_startjahr", "ganz"),
              Feld("rentenbeginn", "datum", pflicht=True), Feld("ende_datum", "datum"),
              Feld("auszahlungsart", standard="rente", hinweis="rente oder einmalzahlung"),
              Feld("auszahlung_kategorie", "kategorie", "auszahlung_kategorie_id", pflicht=True),
              Feld("rente_garantiert_betrag", "zahl"), Feld("rente_voraussichtlich_betrag", "zahl"),
              Feld("einmalzahlung_garantiert_betrag", "zahl"), Feld("einmalzahlung_voraussichtlich_betrag", "zahl"),
              Feld("rueckkaufswert", "zahl"), Feld("auszahlung_bei_tod", "zahl"),
              Feld("auszahlung_inflationsflag", "bool", standard=0), Feld("aktiv", "bool", standard=1)]

    def _payload(self, w):
        _pflicht(w, "bezeichnung", "zahlungskonto", "stand_datum", "rentenbeginn", "auszahlung_kategorie")
        hat = bool(w.get("hat_einzahlung", 1))
        if hat:
            _pflicht(w, "beitrag_kategorie", "beitrag_betrag", "beitrag_startdatum")
            if w.get("beitrag_dynamik_aktiv") and w.get("beitrag_dynamik_prozent") is None:
                raise ZeilenFehler("beitrag_dynamik_aktiv ist gesetzt, aber beitrag_dynamik_prozent fehlt")
        art = (w.get("auszahlungsart") or "rente").lower()
        if art not in ("rente", "einmalzahlung"):
            raise ZeilenFehler("auszahlungsart muss rente oder einmalzahlung sein")
        kw = {k: w.get(k) for k in (
            "anbieter", "beitrag_betrag", "beitrag_monat", "beitrag_startdatum", "beitrag_dynamik_rhythmus_jahre",
            "beitrag_dynamik_prozent", "beitrag_dynamik_startjahr", "ende_datum", "rente_garantiert_betrag",
            "rente_voraussichtlich_betrag", "einmalzahlung_garantiert_betrag", "einmalzahlung_voraussichtlich_betrag",
            "rueckkaufswert", "auszahlung_bei_tod")}
        return _modell(api.AltersvorsorgeIn, bezeichnung=w["bezeichnung"], zahlungskonto_id=w["zahlungskonto"],
                       person_id=w.get("person"), stand_datum=w["stand_datum"], hat_einzahlung=hat,
                       beitrag_kategorie_id=w.get("beitrag_kategorie") if hat else None,
                       beitrag_intervall=(w.get("beitrag_intervall") or "monatlich").lower(),
                       beitrag_tag_des_monats=w.get("beitrag_tag_des_monats") or 1,
                       beitrag_dynamik_aktiv=bool(w.get("beitrag_dynamik_aktiv")) and hat,
                       rentenbeginn=w["rentenbeginn"], auszahlungsart=art,
                       auszahlung_kategorie_id=w["auszahlung_kategorie"],
                       auszahlung_inflationsflag=bool(w.get("auszahlung_inflationsflag")), **kw)

    def anlegen(self, conn, ctx, w):
        try:
            aid = api._av_anlegen(conn, self._payload(w))
        except HTTPException as e:
            raise ZeilenFehler(str(e.detail))
        if w.get("aktiv") == 0:
            conn.execute("UPDATE altersvorsorge SET aktiv = 0 WHERE id = ?", (aid,))
        return aid

    def aendern(self, conn, ctx, id_, w, alt):
        # Ein verschobener Rentenbeginn bereinigt die Phasen wie in der
        # Oberfläche (Punkt 5, v2.0); die Vorschau nennt die Anzahl.
        api._av_aendern(conn, id_, self._payload(w), phasen_bereinigen=True)
        conn.execute("UPDATE altersvorsorge SET aktiv = ? WHERE id = ?", (w.get("aktiv", 1), id_))

    def loeschen(self, conn, ctx, id_, alt):
        api._av_loeschen(conn, id_, inkl_buchungen=False)

    def hinweise_vor_aenderung(self, conn, ctx, id_, w, alt):
        if alt is None or w is None or not w.get("rentenbeginn") or w["rentenbeginn"] == alt["rentenbeginn"]:
            return []
        k = av_modul.phasen_konflikte(conn, alt, date.fromisoformat(w["rentenbeginn"]), api._feiertage_laden(conn))
        teile = []
        if k["loeschen"]:
            teile.append(f"„{alt['bezeichnung']}“: Rentenbeginn verschoben, {len(k['loeschen'])} automatisch erzeugte "
                         "Buchung(en) in der falschen Phase werden gelöscht.")
        if k["angepasst"]:
            teile.append(f"„{alt['bezeichnung']}“: {len(k['angepasst'])} angepasste Buchung(en) im betroffenen "
                         "Zeitraum bleiben stehen – bitte selbst prüfen.")
        return teile


class Wiederkehrend(Bereich):
    key, sheet, tabelle = "wiederkehrend", "Wiederkehrend", "wiederkehrende_buchungen"
    titel_feld = "beschreibung"
    felder = [Feld("konto", "konto", "konto_id", pflicht=True), Feld("person", "person", "person_id"),
              Feld("betrag", "zahl", pflicht=True), Feld("kategorie", "kategorie", "kategorie_id"),
              Feld("zielkonto", "konto", "zielkonto_id", hinweis="gefüllt = Umbuchungs-Vorlage"),
              Feld("sparziel", "sparziel", "sparziel_id", hinweis="nur bei Umbuchung"),
              Feld("beschreibung"), Feld("intervall", pflicht=True, hinweis="monatlich, quartalsweise, halbjaehrlich, jaehrlich"),
              Feld("tag_des_monats", "ganz", pflicht=True), Feld("monat", "ganz"),
              Feld("startdatum", "datum", pflicht=True), Feld("enddatum", "datum"),
              Feld("inflationsflag", "bool", standard=0), Feld("dynamik_aktiv", "bool", standard=0),
              Feld("dynamik_rhythmus_jahre", "ganz"), Feld("dynamik_prozent", "zahl"), Feld("dynamik_startjahr", "ganz"),
              Feld("aktiv", "bool", standard=1)]

    def export_zeilen(self, conn, params):
        return conn.execute("SELECT * FROM wiederkehrende_buchungen ORDER BY beschreibung, id").fetchall()

    def _payload(self, conn, w):
        _pflicht(w, "konto", "betrag", "intervall", "tag_des_monats", "startdatum")
        intervall = w["intervall"].lower()
        if intervall not in ("monatlich", "quartalsweise", "halbjaehrlich", "jaehrlich"):
            raise ZeilenFehler("intervall muss monatlich/quartalsweise/halbjaehrlich/jaehrlich sein")
        umb = bm.umbuchung_kategorie_id(conn)
        if w.get("zielkonto"):
            # Umbuchungs-Vorlage (v1.10): Widersprüche sind Zeilenfehler
            if w.get("kategorie") not in (None, umb):
                raise ZeilenFehler("Bei einer Umbuchung (zielkonto gefüllt) muss kategorie leer sein oder 'Umbuchung' lauten")
            if w.get("inflationsflag") or w.get("dynamik_aktiv"):
                raise ZeilenFehler("Bei einer Umbuchung (zielkonto gefüllt) sind inflationsflag und dynamik_aktiv nicht möglich")
        else:
            if w.get("sparziel"):
                raise ZeilenFehler("sparziel ist nur bei einer Umbuchung (zielkonto gefüllt) möglich")
            if w.get("kategorie") is None:
                raise ZeilenFehler("Kategorie fehlt")
        return _modell(api.WiederkehrendIn, konto_id=w["konto"], betrag=w["betrag"],
                       kategorie_id=w.get("kategorie"), beschreibung=w.get("beschreibung"), intervall=intervall,
                       tag_des_monats=w["tag_des_monats"], monat=w.get("monat"), startdatum=w["startdatum"],
                       enddatum=w.get("enddatum"), inflationsflag=bool(w.get("inflationsflag")),
                       dynamik_aktiv=bool(w.get("dynamik_aktiv")),
                       dynamik_rhythmus_jahre=w.get("dynamik_rhythmus_jahre"),
                       dynamik_prozent=w.get("dynamik_prozent"), dynamik_startjahr=w.get("dynamik_startjahr"),
                       person_id=w.get("person"), zielkonto_id=w.get("zielkonto"), sparziel_id=w.get("sparziel"))

    def anlegen(self, conn, ctx, w):
        vid = api._wk_anlegen(conn, self._payload(conn, w))
        if w.get("aktiv") == 0:
            api._wk_aktiv_setzen(conn, vid, False)
        return vid

    def aendern(self, conn, ctx, id_, w, alt):
        api._wk_aendern(conn, id_, self._payload(conn, w))
        if w.get("aktiv", 1) != alt["aktiv"]:
            api._wk_aktiv_setzen(conn, id_, bool(w.get("aktiv", 1)))

    def loeschen(self, conn, ctx, id_, alt):
        api._wk_loeschen(conn, id_)


class Buchungen(Bereich):
    key, sheet, tabelle = "buchungen", "Buchungen", "buchungen"
    titel_feld = "beschreibung"
    felder = [Feld("datum", "datum", pflicht=True), Feld("konto", "konto", "konto_id", pflicht=True),
              Feld("typ", standard="Ist", hinweis="Ist, Plan oder Uebertragung"),
              Feld("betrag", "zahl", pflicht=True), Feld("kategorie", "kategorie", "kategorie_id"),
              Feld("zielkonto", "konto", "zielkonto_id", hinweis="nur bei Uebertragung"),
              Feld("sparziel", "sparziel", "sparziel_id", hinweis="nur bei Uebertragung von/zum Trägerkonto"),
              Feld("beschreibung"), Feld("herkunft", nur_lesen=True)]

    def export_zeilen(self, conn, params):
        bed, par = [], []
        if params.get("von"):
            bed.append("datum >= ?"); par.append(params["von"])
        if params.get("bis"):
            bed.append("datum <= ?"); par.append(params["bis"])
        where = (" WHERE " + " AND ".join(bed)) if bed else ""
        return conn.execute(f"SELECT * FROM buchungen{where} ORDER BY datum, id", par).fetchall()

    def aus_db(self, conn, row):
        w = super().aus_db(conn, row)
        if row["typ"] == "Uebertragung":
            w["kategorie"] = None   # Übertragungen: Kategorie fest „Umbuchung“, in der Datei leer
        return w

    def nur_lesen_werte(self, conn, ctx, row):
        if row["wiederkehrend_id"]:
            h = "Vorlage"
        elif row["ziel_altersvorsorge_id"]:
            h = "Altersvorsorge"
        elif row["ziel_darlehen_id"]:
            h = "Darlehen"
        else:
            h = "manuell"
        return {"herkunft": h}

    def _pruefen(self, conn, w):
        _pflicht(w, "datum", "konto", "betrag")
        typ = w.get("typ") or ("Uebertragung" if w.get("zielkonto") else "Ist")
        typ = {"ist": "Ist", "plan": "Plan", "uebertragung": "Uebertragung", "übertragung": "Uebertragung"}.get(
            typ.lower(), typ)
        if typ not in ("Ist", "Plan", "Uebertragung"):
            raise ZeilenFehler("typ muss Ist, Plan oder Uebertragung sein")
        if typ == "Uebertragung":
            if not w.get("zielkonto"):
                raise ZeilenFehler("Bei einer Übertragung ist zielkonto Pflicht")
            if w["zielkonto"] == w["konto"]:
                raise ZeilenFehler("zielkonto muss sich vom Konto unterscheiden")
            if w.get("kategorie") not in (None, bm.umbuchung_kategorie_id(conn)):
                raise ZeilenFehler("Bei einer Übertragung muss kategorie leer sein oder 'Umbuchung' lauten")
        else:
            if w.get("zielkonto"):
                raise ZeilenFehler("zielkonto ist nur bei typ Uebertragung möglich")
            if w.get("kategorie") is None:
                raise ZeilenFehler("Kategorie fehlt")
        fehler = bm.sparziel_pruefen(conn, typ, w["konto"], w.get("zielkonto"), w.get("sparziel"))
        if fehler:
            raise ZeilenFehler(fehler)
        return typ

    def anlegen(self, conn, ctx, w):
        typ = self._pruefen(conn, w)
        return bm.buchung_anlegen(conn, konto_id=w["konto"], datum=w["datum"], betrag=w["betrag"],
                                  kategorie_id=w.get("kategorie"), typ=typ, beschreibung=w.get("beschreibung"),
                                  zielkonto_id=w.get("zielkonto"), sparziel_id=w.get("sparziel"))

    def aendern(self, conn, ctx, id_, w, alt):
        typ = self._pruefen(conn, w)
        bm.buchung_bearbeiten(conn, id_, konto_id=w["konto"], datum=w["datum"], betrag=w["betrag"],
                              kategorie_id=w.get("kategorie") or 0, typ=typ, beschreibung=w.get("beschreibung"),
                              zielkonto_id=w.get("zielkonto"), sparziel_id=w.get("sparziel"))

    def loeschen(self, conn, ctx, id_, alt):
        bm.buchung_loeschen(conn, id_)


# Reihenfolge = Ausführungsreihenfolge für Neuanlagen/Änderungen; Löschungen
# laufen danach in umgekehrter Reihenfolge (Abhängige zuerst).
BEREICHE = [Personen(), Konten(), Kategorien(), Regeln(), Budgets(), Darlehen(), Auszahlungen(),
            Altersvorsorge(), Wiederkehrend(), Buchungen()]
BEREICH_NACH_KEY = {b.key: b for b in BEREICHE}
BEREICH_NACH_SHEET = {b.sheet.lower(): b for b in BEREICHE}
# CSV-Bereichsnamen des bisherigen Imports
CSV_ALIAS = {"buchungen": "buchungen", "wiederkehrend": "wiederkehrend", "budget": "budget",
             "darlehen": "darlehen", "altersvorsorge": "altersvorsorge"}


def pruefwert(werte: dict) -> str:
    return hashlib.sha1(json.dumps(werte, sort_keys=True, default=str).encode()).hexdigest()[:16]


# ─── Export ──────────────────────────────────────────────────────────

_KOPF_FILL = PatternFill("solid", start_color="2F7D4F")
_LESEN_FILL = PatternFill("solid", start_color="EEEEEE")
_ANLEITUNG = [
    "Finanzverwaltung – Export / Import (XLSX-Roundtrip, seit Version 2.0)",
    "",
    "Jeder Tab entspricht einem Bereich. Ändern: Werte in der Zeile überschreiben. Neu: Zeile ohne id anfügen.",
    "Löschen: in der Spalte „aktion“ löschen eintragen (Zeilen einfach entfernen löscht NICHTS).",
    "Konten lassen sich per Datei nicht löschen (dafür Administration → Daten zurücksetzen). Personen werden deaktiviert.",
    "Grau hinterlegte Spalten dienen nur zur Information und werden beim Hochladen ignoriert.",
    "Kategorien als vollständiger Pfad („Ausgaben > Wohnen > Strom“), als Name oder Namensteil, solange eindeutig.",
    "Konten, Personen, Sparziele und Darlehen über den Namen oder einen eindeutigen Namensteil. Mehrdeutiges wird als Fehler gemeldet.",
    "Wahrheitswerte: 1 = ja, 0 = nein. Datumsangaben als Datum oder JJJJ-MM-TT.",
    "Die verborgene Spalte _pruefwert nicht verändern: Sie erkennt, ob ein Eintrag nach dem Export in der App geändert wurde.",
    "Beim Hochladen zeigt die App zuerst eine Vorschau. Enthält die Datei einen Fehler, wird nichts übernommen.",
    "Vor der Übernahme legt die App automatisch ein Backup an.",
]


def export_workbook(conn, von: Optional[str] = None, bis: Optional[str] = None, leer: bool = False) -> Workbook:
    """leer=True: nur Kopfzeilen (leere Vorlage), sonst der
    aktuelle Datenbestand (Buchungen optional im Zeitraum von/bis)."""
    ctx = Kontext(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Anleitung"
    for i, text in enumerate(_ANLEITUNG, start=1):
        ws.cell(i, 1, text).font = Font(bold=(i == 1), size=12 if i == 1 else 10)
    if von or bis:
        ws.cell(len(_ANLEITUNG) + 2, 1, f"Buchungen im Export: {von or 'Anfang'} bis {bis or 'Ende'}").font = Font(italic=True)
    ws.column_dimensions["A"].width = 120
    params = {"von": von, "bis": bis}
    for bereich in BEREICHE:
        sh = wb.create_sheet(bereich.sheet)
        kopf = ["id", "aktion"] + [f.name for f in bereich.felder] + ["_pruefwert"]
        sh.append(kopf)
        for idx, name in enumerate(kopf, start=1):
            zelle = sh.cell(1, idx)
            zelle.font = Font(bold=True, color="FFFFFF")
            zelle.fill = _KOPF_FILL
            sh.column_dimensions[get_column_letter(idx)].width = max(12, min(40, len(name) + 4))
        for idx, f in enumerate(bereich.felder, start=3):
            if f.nur_lesen:
                sh.cell(1, idx).comment = Comment("Nur Information – wird beim Hochladen ignoriert.", "Finanzverwaltung")
            elif f.hinweis:
                sh.cell(1, idx).comment = Comment(f.hinweis, "Finanzverwaltung")
        sh.cell(1, 2).comment = Comment("„löschen“ eintragen, um den Eintrag zu löschen.", "Finanzverwaltung")
        letzte = get_column_letter(len(kopf))
        sh.column_dimensions[letzte].hidden = True
        sh.freeze_panes = "C2"
        dv = DataValidation(type="list", formula1='"löschen"', allow_blank=True)
        sh.add_data_validation(dv)
        dv.add("B2:B100000")
        if leer:
            continue
        for row in bereich.export_zeilen(conn, params):
            werte = bereich.aus_db(conn, row)
            lesen = bereich.nur_lesen_werte(conn, ctx, row)
            zeile = [row["id"], None]
            for f in bereich.felder:
                zeile.append(lesen.get(f.name) if f.nur_lesen else _anzeige(ctx, f, werte.get(f.name)))
            zeile.append(pruefwert(werte))
            sh.append(zeile)
            r = sh.max_row
            for idx, f in enumerate(bereich.felder, start=3):
                if f.art == "datum":
                    sh.cell(r, idx).number_format = "YYYY-MM-DD"
                if f.nur_lesen:
                    sh.cell(r, idx).fill = _LESEN_FILL
    return wb


# ─── Einlesen ────────────────────────────────────────────────────────

def _zeilen_xlsx(rohbytes: bytes) -> dict:
    wb = load_workbook(io.BytesIO(rohbytes), data_only=True)
    ergebnis = {}
    for name in wb.sheetnames:
        bereich = BEREICH_NACH_SHEET.get(name.strip().lower())
        if bereich is None:
            continue
        zeilen_roh = list(wb[name].iter_rows(values_only=True))
        if not zeilen_roh:
            continue
        kopf = [_txt(h).lower() for h in zeilen_roh[0]]
        zeilen = []
        for nr, r in enumerate(zeilen_roh[1:], start=2):
            if all(v is None or _txt(v) == "" for v in r):
                continue
            zeilen.append((nr, {kopf[i]: r[i] for i in range(min(len(kopf), len(r))) if kopf[i]}))
        if zeilen:
            ergebnis[bereich.key] = zeilen
    return ergebnis


def _zeilen_csv(rohbytes: bytes, art: str) -> dict:
    if art not in BEREICH_NACH_KEY:
        raise HTTPException(400, "Für CSV-Dateien bitte den Bereich auswählen.")
    text = _text_dekodieren(rohbytes)
    zeilen = []
    for nr, r in enumerate(csv.DictReader(io.StringIO(text)), start=2):
        z = {_txt(k).lower(): v for k, v in r.items() if k is not None}
        if all(_txt(v) == "" for v in z.values()):
            continue
        zeilen.append((nr, z))
    return {art: zeilen} if zeilen else {}


def datei_einlesen(dateiname: str, rohbytes: bytes, art: Optional[str] = None) -> dict:
    if (dateiname or "").lower().endswith((".xlsx", ".xlsm")):
        return _zeilen_xlsx(rohbytes)
    return _zeilen_csv(rohbytes, art or "")


# ─── Verarbeiten (Vorschau und Übernahme) ────────────────────────────

def _titel(bereich, werte, z):
    t = _txt(z.get(bereich.titel_feld))
    return t or (werte or {}).get(bereich.titel_feld) or ""


def _verarbeiten(conn_echt, daten: dict) -> dict:
    """Wendet alle Zeilen in der laufenden Transaktion an und liefert die
    Übersicht. Commit/Rollback macht der Aufrufer."""
    conn = _OhneCommit(conn_echt)
    ctx = Kontext(conn)
    # Explizite Transaktion: sonst würde das RELEASE des äußersten
    # Savepoints bereits committen (SQLite-Verhalten).
    if not conn_echt.in_transaction:
        conn_echt.execute("BEGIN")
    ergebnis = {"bereiche": [], "fehler": 0, "aenderungen": 0, "hinweise": [], "feiertagshinweise": []}
    loeschungen = []   # (bereich, nr, id, alt)
    # Stand zu Beginn: Einträge, die im Lauf als Folge einer anderen Änderung
    # wegfallen (z. B. Phasenbereinigung der Altersvorsorge), werden so
    # erkannt und übersprungen, solange die Datei sie nicht selbst ändert.
    start_werte = {}
    for bereich in BEREICHE:
        for _, z in daten.get(bereich.key) or []:
            try:
                i = int(float(_txt(z.get("id"))))
            except (TypeError, ValueError):
                continue
            row = bereich.zeile_lesen(conn, i)
            if row is not None:
                start_werte[(bereich.key, i)] = bereich.aus_db(conn, row)
    for bereich in BEREICHE:
        zeilen = daten.get(bereich.key)
        if not zeilen:
            continue
        info = {"key": bereich.key, "sheet": bereich.sheet, "neu": 0, "geaendert": 0, "geloescht": 0,
                "unveraendert": 0, "entfallen": 0, "fehler": [], "details": []}
        ergebnis["bereiche"].append(info)
        hat_id_spalte = any("id" in z for _, z in zeilen)
        for nr, z in zeilen:
            id_text = _txt(z.get("id")) if hat_id_spalte else ""
            aktion = _txt(z.get("aktion")).lower()
            conn_echt.execute("SAVEPOINT zeile")
            try:
                if aktion and aktion not in LOESCHEN:
                    raise ZeilenFehler(f"Unbekannte aktion „{aktion}“ (erlaubt: leer oder löschen)")
                id_ = None
                alt = None
                if id_text:
                    try:
                        id_ = int(float(id_text))
                    except ValueError:
                        raise ZeilenFehler(f"Ungültige id „{id_text}“")
                    alt = bereich.zeile_lesen(conn, id_)
                    if alt is None and (bereich.key, id_) in start_werte:
                        vorher = start_werte[(bereich.key, id_)]
                        abweichend = aktion not in LOESCHEN and any(
                            vorher.get(f.name) != _wert_aus_datei(ctx, f, z.get(f.name))
                            for f in bereich.felder if not f.nur_lesen and f.name in z)
                        if abweichend:
                            raise ZeilenFehler("Der Eintrag fällt durch eine andere Änderung dieser Datei weg "
                                               "und kann daher nicht zugleich geändert werden")
                        info["entfallen"] += 1
                        info["details"].append({"zeile": nr, "aktion": "entfällt", "titel": _titel(bereich, None, z),
                                                "felder": ["fällt durch eine andere Änderung dieser Datei weg"]})
                        conn_echt.execute("RELEASE zeile")
                        continue
                    if alt is None:
                        raise ZeilenFehler(f"Eintrag mit id {id_} existiert nicht (mehr)")
                if aktion in LOESCHEN:
                    if id_ is None:
                        raise ZeilenFehler("Löschen ist nur mit id möglich")
                    if not bereich.loeschen_erlaubt:
                        raise ZeilenFehler(f"{bereich.sheet} lassen sich per Datei nicht löschen "
                                           "(dafür Administration → Daten zurücksetzen)")
                    _konflikt_pruefen(bereich, conn, alt, z)
                    loeschungen.append((bereich, info, nr, id_, alt, _titel(bereich, None, z)))
                    conn_echt.execute("RELEASE zeile")
                    continue
                vorhanden = {f.name for f in bereich.felder if not f.nur_lesen and f.name in z}
                werte_datei = {}
                for f in bereich.felder:
                    if f.nur_lesen or f.name not in vorhanden:
                        continue
                    werte_datei[f.name] = _wert_aus_datei(ctx, f, z.get(f.name))
                if alt is None:
                    werte = {f.name: f.standard for f in bereich.felder if not f.nur_lesen}
                    werte.update(werte_datei)
                    neu_id = bereich.anlegen(conn, ctx, werte)
                    ergebnis["feiertagshinweise"] += _feiertagshinweise(conn, bereich, neu_id, werte, nr, z, True)
                    info["neu"] += 1
                    info["details"].append({"zeile": nr, "aktion": "neu", "titel": _titel(bereich, werte, z)})
                else:
                    werte_db = bereich.aus_db(conn, alt)
                    geaendert = {k: (werte_db.get(k), v) for k, v in werte_datei.items() if werte_db.get(k) != v}
                    if not geaendert:
                        info["unveraendert"] += 1
                        conn_echt.execute("RELEASE zeile")
                        continue
                    _konflikt_pruefen(bereich, conn, alt, z)
                    werte = dict(werte_db)
                    werte.update(werte_datei)
                    ergebnis["hinweise"] += bereich.hinweise_vor_aenderung(conn, ctx, id_, werte, alt)
                    bereich.aendern(conn, ctx, id_, werte, alt)
                    ergebnis["feiertagshinweise"] += _feiertagshinweise(conn, bereich, id_, werte, nr, z,
                                                                        "datum" in geaendert)
                    info["geaendert"] += 1
                    felder = {f.name: f for f in bereich.felder}
                    info["details"].append({
                        "zeile": nr, "aktion": "geändert", "titel": _titel(bereich, werte, z),
                        "felder": [f"{k}: {_text(ctx, felder[k], a)} → {_text(ctx, felder[k], n)}"
                                   for k, (a, n) in geaendert.items()]})
                conn_echt.execute("RELEASE zeile")
            except (ZeilenFehler, HTTPException, ValueError, KeyError, TypeError, sqlite3.Error) as e:
                conn_echt.execute("ROLLBACK TO zeile")
                conn_echt.execute("RELEASE zeile")
                text = e.detail if isinstance(e, HTTPException) else str(e)
                info["fehler"].append({"zeile": nr, "titel": _titel(bereich, None, z), "hinweis": text})
                ergebnis["fehler"] += 1
    # Löschungen zuletzt, Abhängige zuerst
    reihenfolge = {b.key: i for i, b in enumerate(BEREICHE)}
    for bereich, info, nr, id_, alt, titel in sorted(loeschungen, key=lambda x: -reihenfolge[x[0].key]):
        conn_echt.execute("SAVEPOINT zeile")
        try:
            ergebnis["hinweise"] += bereich.hinweise_vor_aenderung(conn, ctx, id_, None, alt)
            bereich.loeschen(conn, ctx, id_, alt)
            info["geloescht"] += 1
            info["details"].append({"zeile": nr, "aktion": "gelöscht", "titel": titel})
            conn_echt.execute("RELEASE zeile")
        except (ZeilenFehler, HTTPException, ValueError, sqlite3.Error) as e:
            conn_echt.execute("ROLLBACK TO zeile")
            conn_echt.execute("RELEASE zeile")
            text = e.detail if isinstance(e, HTTPException) else str(e)
            info["fehler"].append({"zeile": nr, "titel": titel, "hinweis": text})
            ergebnis["fehler"] += 1
    ergebnis["hinweise"] = list(dict.fromkeys(ergebnis["hinweise"]))
    for info in ergebnis["bereiche"]:
        info["details"].sort(key=lambda d: d["zeile"])
        info["fehler"].sort(key=lambda d: d["zeile"])
        ergebnis["aenderungen"] += info["neu"] + info["geaendert"] + info["geloescht"]
    return ergebnis


SERIEN_BEREICHE = {"wiederkehrend": "wiederkehrend", "darlehen": "darlehen", "altersvorsorge": "altersvorsorge"}


def _feiertagshinweise(conn, bereich, id_, werte, nr, z, datum_neu) -> list[dict]:
    """Seit v2.6: Feiertagshinweise für neue oder geänderte Zeilen. Einzelne
    Buchungsdaten (Buchungen, Auszahlungen) nur, wenn das Datum neu ist oder
    sich ändert. Serien (Vorlagen, Darlehen, Altersvorsorge) bei jeder
    Änderung, Termine der nächsten 12 Monate wie im Formular."""
    region = feiertage_modul.gewaehlte_region(conn)
    if bereich.key in ("buchungen", "auszahlungen"):
        if not datum_neu or not werte.get("datum"):
            return []
        h = feiertage_modul.datum_hinweis(date.fromisoformat(werte["datum"]), region)
        treffer = [h] if h else []
    elif bereich.key in SERIEN_BEREICHE and id_:
        treffer = api._serien_hinweise(conn, SERIEN_BEREICHE[bereich.key], id_)
    else:
        return []
    titel = _titel(bereich, werte, z)
    return [{"sheet": bereich.sheet, "zeile": nr, "titel": titel, "text": h.get("text") or feiertage_modul.hinweis_text(h)}
            for h in treffer]


def _text(ctx, feld, wert):
    if wert is None or wert == "":
        return "–"
    if feld.art in ("konto", "person", "sparziel", "darlehen", "kategorie"):
        return ctx.name(feld.art, wert) or f"#{wert}"
    return str(wert)


def _konflikt_pruefen(bereich, conn, alt, z):
    datei = _txt(z.get("_pruefwert"))
    if not datei:
        raise ZeilenFehler("Prüfwert fehlt – Änderungen und Löschungen bitte in einer frisch exportierten Datei vornehmen")
    if datei != pruefwert(bereich.aus_db(conn, alt)):
        raise ZeilenFehler("Der Eintrag wurde nach dem Export in der App geändert (Konflikt) – bitte neu exportieren")


def vorschau(conn, dateiname: str, rohbytes: bytes, art: Optional[str]) -> dict:
    daten = datei_einlesen(dateiname, rohbytes, art)
    if not daten:
        return {"bereiche": [], "fehler": 0, "aenderungen": 0, "hinweise": [], "feiertagshinweise": [], "token": None,
                "meldung": "Die Datei enthält keine erkennbaren Tabs bzw. Zeilen."}
    try:
        ergebnis = _verarbeiten(conn, daten)
    finally:
        conn.rollback()
    ergebnis["token"] = None
    if not ergebnis["fehler"] and ergebnis["aenderungen"]:
        jetzt = time.time()
        for t in [t for t, v in _TOKENS.items() if v["bis"] < jetzt]:
            del _TOKENS[t]
        token = secrets.token_urlsafe(16)
        _TOKENS[token] = {"daten": daten, "bis": jetzt + _TOKEN_GUELTIG_SEKUNDEN, "name": dateiname}
        ergebnis["token"] = token
    return ergebnis


def ausfuehren(conn, token: Optional[str]) -> dict:
    eintrag = _TOKENS.pop(token or "", None)
    if eintrag is None or eintrag["bis"] < time.time():
        raise HTTPException(422, "Die Vorschau ist abgelaufen. Bitte die Datei erneut hochladen.")
    datei = eintrag["name"] or "Datei"
    sicherung = backup.backup_erstellen(conn, "manuell", beschreibung=f"vor Import von: {datei}")
    try:
        ergebnis = _verarbeiten(conn, eintrag["daten"])
        if ergebnis["fehler"]:
            conn.rollback()
            ergebnis["uebernommen"] = False
            ergebnis["meldung"] = ("Seit der Vorschau hat sich der Datenbestand geändert – es wurde nichts "
                                   "übernommen. Bitte die Datei erneut hochladen.")
            return ergebnis
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    # Fällige Termine neu angelegter/geänderter Quellen buchen (wie beim Programmstart)
    feiertage = api._feiertage_laden(conn)
    wk_modul.erzeuge_faellige_ist_buchungen(conn, feiertage)
    from .core import darlehensraten
    darlehensraten.erzeuge_faellige_ist_buchungen(conn, feiertage)
    av_modul.erzeuge_faellige_ist_buchungen(conn, feiertage)
    ergebnis["uebernommen"] = True
    ergebnis["backup"] = sicherung
    return ergebnis
